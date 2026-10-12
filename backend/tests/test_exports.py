"""Downloads: the Word document and Markdown, built from a saved version or the current state."""

from __future__ import annotations

import asyncio
import io
import json
import zipfile

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn
from google.adk.artifacts import InMemoryArtifactService
from google.adk.sessions import InMemorySessionService
from google.genai import types

from translator import exports
from translator.document import state as st
from translator.exports.docx import to_docx


def _png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), False)
    pix.clear_with(200)
    return pix.tobytes("png")


SEGS = [
    {"id": "p1.1", "page": 1, "kind": "heading", "source": "## Ergebnis", "target": "## Result", "status": "reviewed"},
    {"id": "p1.2", "page": 1, "kind": "paragraph", "source": "Die **Anlage** ist abgenommen.", "target": "The **plant** is accepted.", "status": "confirmed"},
    {"id": "p1.3", "page": 1, "kind": "list", "source": "- eins\n- zwei", "target": "- one\n- two", "status": "draft"},
    {"id": "p1.4", "page": 1, "kind": "table", "source": "| Nr. | Frist |\n|---|---|\n| 1 | 30.04. |", "target": "| No. | Deadline |\n|---|---|\n| 1 | 30.04. |", "status": "draft"},
    {"id": "p2.1", "page": 2, "kind": "figure", "source": "Schaltschrank", "target": "Control cabinet", "image": "figure-p2-1.png", "status": "draft"},
]


def test_docx_keeps_structure_language_and_figures():
    doc = Document(io.BytesIO(to_docx(SEGS, "target", "en", {"figure-p2-1.png": _png()})))
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Heading 2", "Result") in styles
    assert ("List Bullet", "one") in styles and ("List Bullet", "two") in styles
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "plant" in bold
    assert doc.tables[0].cell(0, 1).text == "Deadline" and doc.tables[0].cell(1, 1).text == "30.04."
    assert len(doc.inline_shapes) == 1 and ("Caption", "Control cabinet") in styles
    lang = doc.styles["Normal"].element.rPr.find(qn("w:lang"))
    assert lang.get(qn("w:val")) == "en"


def test_docx_for_chinese_sets_east_asian_font():
    doc = Document(io.BytesIO(to_docx(SEGS[:2], "target", "zh-Hans", {})))
    rpr = doc.styles["Normal"].element.rPr
    assert rpr.find(qn("w:rFonts")).get(qn("w:eastAsia")) == "SimSun"
    assert rpr.find(qn("w:lang")).get(qn("w:eastAsia")) == "zh-Hans"


def test_export_build_uses_the_version_snapshot():
    async def main():
        sessions, artifacts = InMemorySessionService(), InMemoryArtifactService()
        seg = {"id": "p1.1", "page": 1, "kind": "paragraph", "source": "Die Anlage ist abgenommen und freigegeben.",
               "target": "Current text (not saved as a version).", "translatable": True}
        s = await sessions.create_session(app_name="t", user_id="u", state={
            "doc:filename": "a.pdf", "doc:source_artifact": "source.pdf", "doc:target_lang": "English",
            "doc:seg_ids": ["p1.1"], st.seg_key("p1.1"): seg,
            "doc:versions": [{"version": 1, "note": "Draft", "changed": ["p1.1"], "artifact_version": 0}]})
        await artifacts.save_artifact(app_name="t", user_id="u", session_id=s.id, filename=st.TRANSLATION_ARTIFACT,
                                      artifact=types.Part(text=json.dumps({"p1.1": "The plant is accepted and released."})))
        assert [f["id"] for f in exports.available()] == ["docx", "md"]

        content, media, name = await exports.build(artifacts, s, "docx", 1)
        assert media.endswith("wordprocessingml.document") and name == "a.English-v1.docx"
        assert "accepted and released" in "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)  # version 1

        md, media, name = await exports.build(artifacts, s, "md", None)
        assert media.startswith("text/markdown") and b"Current text" in md  # no version → current state

    asyncio.run(main())


def test_markdown_export_with_figures_is_a_zip():
    async def main():
        sessions, artifacts = InMemorySessionService(), InMemoryArtifactService()
        s = await sessions.create_session(app_name="t", user_id="u", state={
            "doc:filename": "b.pdf", "doc:target_lang": "English", "doc:seg_ids": ["p2.1"], st.seg_key("p2.1"): {**SEGS[-1], "translatable": True},
            "doc:versions": []})
        await artifacts.save_artifact(app_name="t", user_id="u", session_id=s.id, filename="figure-p2-1.png",
                                      artifact=types.Part.from_bytes(data=_png(), mime_type="image/png"))
        content, media, name = await exports.build(artifacts, s, "md", None)
        assert media == "application/zip" and name.endswith(".zip")
        assert sorted(zipfile.ZipFile(io.BytesIO(content)).namelist()) == ["b.English.md", "figures/figure-p2-1.png"]

    asyncio.run(main())


def _xml_text(part) -> str:
    """Header/footer text with Word fields shown as {PAGE} / {NUMPAGES}."""
    out = []
    for el in part._element.iter(qn("w:t"), qn("w:fldSimple")):
        if el.tag == qn("w:fldSimple"):
            out.append("{" + el.get(qn("w:instr")).split()[0] + "}")
        elif el.getparent().getparent().tag != qn("w:fldSimple"):
            out.append(el.text or "")
    return "".join(out)


def test_docx_puts_page_furniture_in_header_and_footer():
    segs = [
        {"id": "p1.1", "page": 1, "kind": "header", "source": "Vertraulich – UW-2291", "target": "Confidential – UW-2291", "status": "draft"},
        *SEGS[:2],
        {"id": "p1.9", "page": 1, "kind": "footer", "source": "Seite 1 von 2", "target": "Page 1 of 2", "status": "draft"},
        {"id": "p2.1", "page": 2, "kind": "header", "source": "Vertraulich – UW-2291", "target": "Confidential – UW-2291", "status": "draft"},
        {"id": "p2.2", "page": 2, "kind": "paragraph", "source": "Ende.", "target": "End.", "status": "draft"},
        {"id": "p2.9", "page": 2, "kind": "footer", "source": "Seite 2 von 2", "target": "Page 2 of 2", "status": "draft"},
    ]
    doc = Document(io.BytesIO(to_docx(segs, "target", "en", {})))
    # Word repeats these on every page: the page number must be a field, not page 1's "1".
    assert _xml_text(doc.sections[0].header) == "Confidential – UW-2291"  # constant numbers stay text
    assert _xml_text(doc.sections[0].footer) == "Page {PAGE} of {NUMPAGES}"
    assert all(p.text not in ("Confidential – UW-2291", "Page 1 of 2") for p in doc.paragraphs)  # not in the body


@pytest.mark.parametrize("footers, expected", [
    ({2: "- 2 -", 3: "- 3 -"}, "- {PAGE} -"),  # no footer on the title page
    ({1: "第 1 页 共 3 页", 2: "第 2 页 共 3 页", 3: "第 3 页 共 3 页"}, "第 {PAGE} 页 共 {NUMPAGES} 页"),
    ({3: "Page 3 of 3"}, "Page {PAGE} of {NUMPAGES}"),  # one sample, but clearly a page number
    ({1: "Revision 3, 2026"}, "Revision 3, 2026"),  # numbers that aren't page numbers stay text
    ({1: "Projekt 2 · Blatt 1", 2: "Projekt 2 · Blatt 2", 3: "Projekt 2 · Blatt 3"}, "Projekt 2 · Blatt {PAGE}"),
])
def test_running_footer_page_numbers_become_fields(footers, expected):
    segs = [{"id": f"p{p}.1", "page": p, "kind": "paragraph", "source": "x", "target": "x"} for p in (1, 2, 3)]
    segs += [{"id": f"p{p}.9", "page": p, "kind": "footer", "source": t, "target": t} for p, t in footers.items()]
    doc = Document(io.BytesIO(to_docx(sorted(segs, key=lambda s: s["page"]), "target", "en", {})))
    assert _xml_text(doc.sections[0].footer) == expected
