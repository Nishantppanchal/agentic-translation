"""Office files built in code for the filter tests (no binary fixtures in the repo)."""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

# Paragraphs python-docx can't create: hyperlink, field (with a translatable result), content control,
# tracked changes, a text box stored twice (mc:Choice + mc:Fallback), and a footnote reference.
EXTRA_BODY = """
<w:p><w:r><w:t xml:space="preserve">Siehe </w:t></w:r><w:hyperlink r:id="rIdLink1"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr><w:t>Website des Netzbetreibers</w:t></w:r></w:hyperlink><w:r><w:t xml:space="preserve"> für Details.</w:t></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Details in Abschnitt </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> REF _Ref1 \\h </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Ergebnis der Prüfung</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r><w:r><w:t>.</w:t></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Prüfer: </w:t></w:r><w:sdt><w:sdtPr><w:alias w:val="Prüfer"/><w:tag w:val="inspector"/></w:sdtPr><w:sdtContent><w:r><w:t>Name eintragen</w:t></w:r></w:sdtContent></w:sdt></w:p>
<w:p><w:r><w:t xml:space="preserve">Die Anlage ist </w:t></w:r><w:del w:id="1" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>nicht </w:delText></w:r></w:del><w:ins w:id="2" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:t xml:space="preserve">vollständig </w:t></w:r></w:ins><w:r><w:t>geerdet.</w:t></w:r></w:p>
<w:p><w:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"><mc:Choice Requires="wps"><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:txbx><w:txbxContent><w:p><w:r><w:t>Hinweis im Textfeld</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox><w:txbxContent><w:p><w:r><w:t>Hinweis im Textfeld</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Die Frist ist verbindlich.</w:t></w:r><w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="1"/></w:r></w:p>
"""

FOOTNOTES = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes {W_NS}>
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Gemäß VOB/B § 12.</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""


def docx_full() -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Vertraulich – Projekt UW-2291"
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Seite "
    doc.add_heading("Abnahmeprotokoll", level=1)
    p = doc.add_paragraph("Die ")
    p.add_run("Kabelpritsche").bold = True
    p.add_run(" im Kellergeschoss ist ")
    p.add_run("nicht").italic = True
    p.add_run(" geerdet.\tFrist: 30.04.2026")
    table = doc.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "Nr.", (0, 1): "Beschreibung", (1, 0): "1", (1, 1): "Beschriftung unvollständig"}.items():
        table.cell(r, c).text = text
    buf = io.BytesIO()
    doc.save(buf)
    return _inject(buf.getvalue())


def _inject(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "word/document.xml":
                xml = content.decode()
                xml = xml.replace("<w:sectPr", EXTRA_BODY + "<w:sectPr", 1)
                content = xml.encode()
            elif info.filename == "word/_rels/document.xml.rels":
                xml = content.decode().replace(
                    "</Relationships>",
                    '<Relationship Id="rIdLink1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                    'Target="https://example.com/netz" TargetMode="External"/>'
                    '<Relationship Id="rIdFn1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" '
                    'Target="footnotes.xml"/></Relationships>',
                )
                content = xml.encode()
            elif info.filename == "[Content_Types].xml":
                xml = content.decode().replace(
                    "</Types>",
                    '<Override PartName="/word/footnotes.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
                )
                content = xml.encode()
            elif re.fullmatch(r"word/footer\d*\.xml", info.filename):
                # A PAGE field in the footer: "Seite <PAGE>"
                xml = content.decode()
                xml = xml.replace(
                    "<w:t xml:space=\"preserve\">Seite </w:t></w:r>",
                    "<w:t xml:space=\"preserve\">Seite </w:t></w:r><w:r><w:fldChar w:fldCharType=\"begin\"/></w:r>"
                    "<w:r><w:instrText xml:space=\"preserve\"> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType=\"separate\"/></w:r>"
                    "<w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType=\"end\"/></w:r>",
                )
                content = xml.encode()
            z.writestr(info, content)
        z.writestr("word/footnotes.xml", FOOTNOTES)
    return out.getvalue()


def pptx_deck() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Abnahme Umspannwerk Nord"
    tf = slide.placeholders[1].text_frame
    tf.text = "Erste Zeile"
    first = tf.paragraphs[0]
    first._p.append(first._p.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}br", {}))
    first.add_run().text = "nach dem Umbruch"
    p = tf.add_paragraph()
    bold = p.add_run()
    bold.text = "Wichtig: "
    bold.font.bold = True
    p.add_run().text = "Erdung prüfen"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4.5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Frist"
    table.cell(1, 0).text = "30.04.2026"
    slide.notes_slide.notes_text_frame.text = "Sprechernotiz zur Abnahme"
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Ergebnis"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xlsx_book() -> bytes:
    """xlsxwriter writes shared strings, as Excel does."""
    import xlsxwriter

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Mängel")
    bold = wb.add_format({"bold": True})
    ws.write("A1", "Beschreibung")
    ws.write("B1", "Frist")
    ws.write("A2", "Kabelpritsche nicht geerdet")
    ws.write("B2", "2026")
    ws.write_rich_string("A3", bold, "Wichtig:", " Erdung prüfen")
    ws.write_formula("C2", "=LEN(A2)")
    ws.write("A4", "Beschreibung")
    ws.write_number("B4", 42)
    notes = wb.add_worksheet("Notizen")
    notes.write("A1", "Nachbegehung geplant")
    wb.close()
    return buf.getvalue()
