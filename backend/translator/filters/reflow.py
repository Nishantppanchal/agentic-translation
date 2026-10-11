"""Reflowed exports built from segments: DOCX, a bilingual review DOCX, and standalone HTML.

"Reflowed" = structure (headings, lists, tables, figures) is kept, the original page layout is not;
use `pdf_layout` for a look-alike PDF.
"""

from __future__ import annotations

import base64
import html
import io
import re
from collections import Counter

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

from ..languages import EAST_ASIAN_FONTS, RTL
from . import inline
from .align import markdown_grid

_md = MarkdownIt("commonmark").enable("table")
_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_DIGITS = re.compile(r"\d+")
STATUS_SHADE = {"new": "F1F1F1", "draft": "FFF4D6", "confirmed": "E6F4EA", "reviewed": "E8EAFD"}


def _text(seg: dict, field: str) -> str:
    text = (seg.get(field) if field == "source" else seg.get("target")) or seg.get("source") or ""
    return inline.strip(text)  # Office/XLIFF formatting codes have no place in a reflowed document


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #


def _set_language(doc: Document, lang: str) -> None:
    """Document language (spell-check, hyphenation) and an East Asian font for CJK targets."""
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    if lang in EAST_ASIAN_FONTS:
        lang_el.set(qn("w:eastAsia"), lang)
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), EAST_ASIAN_FONTS[lang])
    if lang.split("-")[0] in RTL:
        bidi = OxmlElement("w:rtl")
        rpr.append(bidi)


def _inline(paragraph, text: str, bold: bool = False) -> None:
    """Markdown inline (bold, italic, code, links) → runs."""
    tokens = _md.parseInline(text)
    children = tokens[0].children if tokens else []
    state = {"strong": bold, "em": False}
    for t in children or []:
        if t.type == "strong_open":
            state["strong"] = True
        elif t.type == "strong_close":
            state["strong"] = bold
        elif t.type == "em_open":
            state["em"] = True
        elif t.type == "em_close":
            state["em"] = False
        elif t.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif t.type in ("text", "code_inline"):
            run = paragraph.add_run(t.content)
            run.bold, run.italic = state["strong"], state["em"]
            if t.type == "code_inline":
                run.font.name = "Consolas"


def _running_parts(segments: list[dict], kind: str, field: str) -> list[str | tuple[str, str]]:
    """The running header/footer as literal text and Word fields: ("PAGE", "1") / ("NUMPAGES", "2").

    Word repeats one header/footer on every page, so a page number copied from page 1 would print
    "Page 1 of 2" on page 2 as well. Numbers that follow the page number become PAGE fields, and a
    constant number next to them that equals the page count becomes NUMPAGES.
    """
    by_page: dict[int, str] = {}
    for s in segments:
        text = _text(s, field).strip() if s.get("kind") == kind else ""
        if text:
            by_page[s["page"]] = f"{by_page[s['page']]} · {text}" if s["page"] in by_page else text
    if not by_page:
        return []
    total = max(s.get("page") or 1 for s in segments)
    shape = lambda t: _DIGITS.sub("#", t)  # noqa: E731
    common = Counter(shape(t) for t in by_page.values()).most_common(1)[0][0]
    pages = sorted(p for p, t in by_page.items() if shape(t) == common)
    first = by_page[pages[0]]
    columns = list(zip(*(map(int, _DIGITS.findall(by_page[p])) for p in pages)))
    # With a single sample, only a short page-number line ("Seite 1 von 2", "- 2 -") counts.
    page_like = len(pages) > 1 or len(re.findall(r"\w+", _DIGITS.sub("", first))) <= 3
    page_col = next((i for i, col in enumerate(columns) if page_like and list(col) == pages), None)  # one per line
    kinds: list[str | None] = [("PAGE" if i == page_col else None) for i in range(len(columns))]
    if page_col is not None:
        kinds = [k or ("NUMPAGES" if set(col) == {total} else None) for k, col in zip(kinds, columns)]
    parts: list[str | tuple[str, str]] = []
    pos = 0
    for m, k in zip(_DIGITS.finditer(first), kinds):
        if k:
            parts += [first[pos:m.start()], (k, m.group())]
            pos = m.end()
    parts.append(first[pos:])
    return [p for p in parts if p]


def _add_field(paragraph, instr: str, shown: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} \\* MERGEFORMAT ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = shown  # what Word shows until it updates the field
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def _table(doc: Document, grid: list[list[str]], style: str = "Table Grid") -> None:
    if not grid:
        return
    cols = max(len(r) for r in grid)
    table = doc.add_table(rows=len(grid), cols=cols)
    table.style = style
    for r, row in enumerate(grid):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = ""
            _inline(cell.paragraphs[0], row[c] if c < len(row) else "", bold=(r == 0))


def to_docx(segments: list[dict], field: str, lang: str, figures: dict[str, bytes]) -> bytes:
    doc = Document()
    _set_language(doc, lang)
    # Running header/footer → the document's real header/footer, with live page-number fields.
    section = doc.sections[0]
    for kind, part in (("header", section.header), ("footer", section.footer)):
        paragraph = part.paragraphs[0]
        for item in _running_parts(segments, kind, field):
            if isinstance(item, tuple):
                _add_field(paragraph, *item)
            else:
                paragraph.add_run(item)
    for seg in segments:
        text = _text(seg, field)
        kind = seg.get("kind")
        if kind in ("header", "footer"):
            continue
        if kind == "figure":
            image = figures.get(seg.get("image") or "")
            if image:
                doc.add_picture(io.BytesIO(image), width=Inches(5.5))
            if text.strip():
                cap = doc.add_paragraph(style="Caption")
                cap.add_run(text.replace("\n", " · "))
            continue
        if kind == "table":
            _table(doc, markdown_grid(text))
            doc.add_paragraph()
            continue
        if kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(text.strip("`\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        for line in text.splitlines() or [""]:
            if not line.strip():
                continue
            if m := _HEADING.match(line):
                doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            elif m := _BULLET.match(line):
                _inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            elif m := _NUMBERED.match(line):
                _inline(doc.add_paragraph(style="List Number"), m.group(1))
            else:
                _inline(doc.add_paragraph(), line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def to_bilingual_docx(segments: list[dict], meta: dict, notes: dict[str, list[str]]) -> bytes:
    """Review table (landscape): ID | Source | Target | Status | Notes — the classic LSP review format."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.6))
    _set_language(doc, meta.get("target_code", "en"))
    title = doc.add_paragraph()
    title.add_run(f"{meta.get('filename', 'Document')} — {meta.get('source_lang')} → {meta.get('target_lang')}").bold = True
    doc.add_paragraph(f"Version {meta.get('version', '?')} · {len(segments)} segments").runs[0].font.size = Pt(9)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    header = table.rows[0]
    trpr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)  # repeat the header row on every page
    for cell, label in zip(header.cells, ["ID", "Source", "Target", "Status", "Notes"]):
        cell.text = ""
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        _shade(cell, "D9D9D9")
    widths = [Inches(0.6), Inches(3.9), Inches(3.9), Inches(0.9), Inches(1.6)]
    for seg in segments:
        if seg.get("kind") == "figure" and not seg.get("source"):
            continue
        row = table.add_row().cells
        status = seg.get("status", "draft")
        values = [
            seg["id"],
            inline.display(seg.get("source")),
            inline.display(seg.get("target")),
            status + (" 🔒" if seg.get("locked") else ""),
            "\n".join(notes.get(seg["id"], [])),
        ]
        for i, (cell, value) in enumerate(zip(row, values)):
            cell.width = widths[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(9 if i in (1, 2) else 8)
            if i == 4:
                run.font.color.rgb = RGBColor(0x8A, 0x4B, 0x00)
        _shade(row[3], STATUS_SHADE.get(status, "FFFFFF"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #

_CSS = """
body{font:16px/1.6 system-ui,-apple-system,"Segoe UI",sans-serif;max-width:46rem;margin:3rem auto;padding:0 1rem;color:#1f1d1a}
table{border-collapse:collapse;margin:1rem 0}th,td{border:1px solid #ccc;padding:.3rem .6rem;text-align:left}th{background:#f2f2f2}
figure{margin:1.5rem 0}figure img{max-width:100%;border:1px solid #ddd}figcaption{font-size:.85rem;color:#555;white-space:pre-line}
hr.page{border:0;border-top:1px dashed #ccc;margin:2rem 0}
"""


def to_html(segments: list[dict], field: str, lang: str, title: str, figures: dict[str, bytes]) -> str:
    parts, page = [], None
    for seg in segments:
        if seg.get("kind") in ("header", "footer"):
            continue  # page furniture has no place in a reflowed web page
        if page is not None and seg.get("page") != page:
            parts.append('<hr class="page">')
        page = seg.get("page")
        text = _text(seg, field)
        if seg.get("kind") == "figure":
            image = figures.get(seg.get("image") or "")
            img = f'<img src="data:image/png;base64,{base64.b64encode(image).decode()}" alt="">' if image else ""
            parts.append(f"<figure>{img}<figcaption>{html.escape(text)}</figcaption></figure>")
        else:
            parts.append(_md.render(text))
    direction = "rtl" if lang.split("-")[0] in RTL else "ltr"
    return (
        f'<!doctype html><html lang="{lang}" dir="{direction}"><head><meta charset="utf-8">'
        f"<title>{html.escape(title)}</title><style>{_CSS}</style></head><body>{''.join(parts)}</body></html>"
    )
