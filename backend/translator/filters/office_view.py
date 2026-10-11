"""Excel and PowerPoint files as they are laid out, for the Document tab: sheets as grids, slides as canvases.

Built from the original file (its source text never changes). Every text cell or text paragraph carries
the id of its segment, so the UI shows the source or any version's translation in place; numbers, dates,
formula results and pictures come from the file. Not a renderer — no charts, themes or conditional
formats — just enough structure to read the file the way it is arranged.
"""

from __future__ import annotations

import base64
import datetime as dt
import io
import re

from . import inline
from .ooxml import in_fallback

MAX_ROWS, MAX_COLS = 300, 40  # a sheet larger than this is shown in part (the file itself is complete)
MAX_IMAGE = 600_000  # bytes; bigger pictures are left out of the slide view
EMU_PER_PT = 12700


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #


def _plain(source: str) -> str:
    return " ".join(inline.strip(source or "").split())


def _decimals(fmt: str) -> int:
    m = re.search(r"\.(0+)", fmt)
    return len(m.group(1)) if m else 0


def _number(value: float, fmt: str) -> str:
    """A number roughly as Excel shows it: percentages, thousands separators, fixed decimals."""
    fmt = (fmt or "General").split(";")[0]
    if "%" in fmt:
        return f"{value * 100:.{_decimals(fmt)}f}%"
    if fmt == "General":
        return f"{value:g}" if isinstance(value, float) else str(value)
    sep = "," if "#,##" in fmt else ""
    return f"{value:{sep}.{_decimals(fmt)}f}"


def xlsx_view(data: bytes, segments: list[dict]) -> dict:
    from openpyxl import load_workbook

    by_text = {}
    for s in segments:  # a text cell → its segment (shared strings are one segment however often they're used)
        by_text.setdefault(_plain(s["source"]), s["id"])
    formulas = load_workbook(io.BytesIO(data))  # formula text
    values = load_workbook(io.BytesIO(data), data_only=True)  # cached results
    sheets = []
    for ws_f in formulas.worksheets:
        ws = values[ws_f.title]
        rows, cols = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
        cells = []
        for row in ws.iter_rows(min_row=1, max_row=rows, max_col=cols):
            for c in row:
                if c.value is None:
                    continue
                f = ws_f.cell(row=c.row, column=c.column).value
                cell: dict = {"r": c.row, "c": c.column, "ref": c.coordinate}
                if isinstance(f, str) and f.startswith("="):
                    cell["formula"] = f
                if isinstance(c.value, str):
                    cell.update(kind="text", value=c.value, seg=by_text.get(_plain(c.value)))
                elif isinstance(c.value, bool):
                    cell.update(kind="bool", value="TRUE" if c.value else "FALSE")
                elif isinstance(c.value, (dt.datetime, dt.date)):
                    cell.update(kind="date", value=c.value.strftime("%Y-%m-%d"))
                elif isinstance(c.value, (int, float)):
                    cell.update(kind="number", value=_number(c.value, c.number_format))
                else:
                    cell.update(kind="text", value=str(c.value))
                if c.font is not None and c.font.b:
                    cell["bold"] = True
                fill = c.fill.fgColor.rgb if c.fill is not None and c.fill.fill_type == "solid" else None
                if isinstance(fill, str) and len(fill) == 8 and fill not in ("00000000", "FFFFFFFF"):
                    cell["fill"] = "#" + fill[2:]
                if c.alignment is not None and c.alignment.horizontal in ("center", "right", "left"):
                    cell["align"] = c.alignment.horizontal
                if c.alignment is not None and c.alignment.wrap_text:
                    cell["wrap"] = True
                cells.append(cell)
        widths = {}
        for letter, dim in ws_f.column_dimensions.items():
            if dim.width:
                for col in range(dim.min or 0, (dim.max or 0) + 1):
                    widths[col] = round(dim.width * 7 + 5)  # Excel width (characters) → px
        merged = [[m.min_row, m.min_col, m.max_row, m.max_col] for m in ws_f.merged_cells.ranges]
        sheets.append({
            "name": ws_f.title, "rows": rows, "cols": cols, "cells": cells, "widths": widths, "merged": merged,
            "truncated": ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS,
        })
    return {"format": "xlsx", "sheets": sheets}


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_INDEX = "data-seg-index"  # set on each a:p while reading (in memory only): its position in the part


def _number_paragraphs(root) -> None:
    """Number a part's paragraphs the way the extractor does (`office._paragraphs`), so each maps to its
    segment's `loc` (`part#n`). Copies inside mc:Fallback aren't counted."""
    n = 0
    for el in root.iter(f"{_A}p"):
        if in_fallback(el):
            continue
        el.set(_INDEX, str(n))
        n += 1


def _paragraphs(text_frame, part: str, by_loc: dict[str, str], default_size: float | None) -> list[dict]:
    out = []
    for p in text_frame.paragraphs:
        n = p._p.get(_INDEX)
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        out.append({
            "seg": by_loc.get(f"{part}#{n}") if n is not None else None,
            "text": "".join(r.text for r in p.runs) or p.text,
            "level": p.level,
            "bold": any(r.font.bold for r in p.runs),
            "size": max(sizes) if sizes else default_size,
            "align": {1: "left", 2: "center", 3: "right"}.get(int(p.alignment)) if p.alignment is not None else None,
        })
    return out


def _fill(shape) -> str | None:
    try:
        if shape.fill.type == 1:  # MSO_FILL.SOLID
            return "#" + str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError, NotImplementedError):
        pass
    return None


def _shapes(shapes, part: str, by_loc: dict[str, str], place) -> list[dict]:
    """Shapes on a slide (groups flattened) with their box as fractions of the slide."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.find(f"{_A}xfrm")
            if xfrm is None:
                out += _shapes(shape.shapes, part, by_loc, place)
                continue
            off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
            ch_off, ch_ext = xfrm.find(f"{_A}chOff"), xfrm.find(f"{_A}chExt")
            sx = int(ext.get("cx")) / max(1, int(ch_ext.get("cx")))
            sy = int(ext.get("cy")) / max(1, int(ch_ext.get("cy")))

            def inner(x, y, w, h, off=off, ch_off=ch_off, sx=sx, sy=sy):  # child space → slide space
                return place(int(off.get("x")) + (x - int(ch_off.get("x"))) * sx,
                             int(off.get("y")) + (y - int(ch_off.get("y"))) * sy, w * sx, h * sy)

            out += _shapes(shape.shapes, part, by_loc, inner)
            continue
        if shape.left is None or shape.width is None:
            continue
        box = place(shape.left, shape.top, shape.width, shape.height)
        ph = None
        if shape.is_placeholder:
            try:
                kind = shape.placeholder_format.type
                ph = getattr(kind, "name", str(kind)).lower()
            except ValueError:
                ph = None
        default = 40.0 if ph in ("title", "center_title") else 32.0 if ph == "subtitle" else 24.0 if ph in ("body", "object") else 18.0
        item: dict = {**box, "placeholder": ph, "fill": _fill(shape)}
        if shape.has_text_frame:
            item.update(kind="text", paragraphs=_paragraphs(shape.text_frame, part, by_loc, default),
                        bullets=ph in ("body", "object"))
        elif getattr(shape, "has_table", False) and shape.has_table:
            item.update(kind="table", rows=[[_paragraphs(cell.text_frame, part, by_loc, 14.0) for cell in row.cells]
                                            for row in shape.table.rows])
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blob = shape.image.blob
            if len(blob) > MAX_IMAGE:
                continue
            item.update(kind="picture", image=f"data:{shape.image.content_type};base64,{base64.b64encode(blob).decode()}")
        else:
            continue
        out.append(item)
    return out


def pptx_view(data: bytes, segments: list[dict]) -> dict:
    from pptx import Presentation

    by_loc = {s["loc"]: s["id"] for s in segments if s.get("loc")}
    prs = Presentation(io.BytesIO(data))
    width, height = prs.slide_width, prs.slide_height

    def place(x, y, w, h):
        return {"x": x / width, "y": y / height, "w": w / width, "h": h / height}

    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        part = str(slide.part.partname).lstrip("/")
        _number_paragraphs(slide._element)
        item = {"number": number, "shapes": _shapes(slide.shapes, part, by_loc, place), "notes": []}
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide
            _number_paragraphs(notes._element)
            item["notes"] = _paragraphs(notes.notes_text_frame, str(notes.part.partname).lstrip("/"), by_loc, 12.0)
        slides.append(item)
    return {"format": "pptx", "width_pt": width / EMU_PER_PT, "height_pt": height / EMU_PER_PT, "slides": slides}


def view(data: bytes, fmt: str, segments: list[dict]) -> dict | None:
    if fmt == "xlsx":
        return xlsx_view(data, segments)
    if fmt == "pptx":
        return pptx_view(data, segments)
    return None
