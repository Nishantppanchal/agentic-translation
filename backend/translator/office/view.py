"""Excel sheets as grids for the Layout view: what each cell shows, roughly formatted (bold, italic, fill, colour,
alignment, wrap, borders), column widths, row heights and merged cells — and which segment a text cell is, so the
view shows the source or the translation in place. Read with openpyxl (values as last calculated). Charts and
pictures are listed, not drawn.
"""

from __future__ import annotations

import datetime as dt
import io

import openpyxl
from openpyxl.utils import get_column_letter

from . import markup
from .sheets import Book, _cells, read

MAX_ROWS, MAX_COLS = 300, 40


def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb not in ("00000000",):
            return "#" + c.rgb[2:]
    except Exception:  # noqa: BLE001 — theme / indexed colours: left to the view's defaults
        pass
    return None


def _value(v, fmt: str) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d") if not isinstance(v, dt.datetime) or not v.hour and not v.minute else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float):
        if "%" in (fmt or ""):
            return f"{v * 100:.{(fmt.split('.')[1].count('0') if '.' in fmt else 0)}f}%"
        return f"{v:,.2f}".rstrip("0").rstrip(".") if "," in (fmt or "") else (f"{v:g}")
    return str(v)


def sheet_grid(data: bytes, segments: list[dict]) -> dict:
    """{"sheets": [{name, rows: [[cell…]], widths, heights, merges, objects, truncated}]}; a text cell carries the
    id of its segment (`seg`), so the view shows it from the document state."""
    by_source = {s["source"]: s["id"] for s in segments}
    book = Book(data)
    strings = book.strings()
    seg_of: dict[tuple[str, str], str] = {}
    for _i, name, ref, kind, value, _root in _cells(book):
        pieces = read(strings[value]) if kind == "s" and value < len(strings) else read(value) if kind == "inline" else []
        if (sid := by_source.get(markup.encode(pieces).text)) is not None:
            seg_of[(name, ref)] = sid
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, rich_text=False)
    out = []
    for ws in wb.worksheets:
        rows_n, cols_n = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
        rows = []
        for r in range(1, rows_n + 1):
            row = []
            for c in range(1, cols_n + 1):
                cell = ws.cell(r, c)
                ref = f"{get_column_letter(c)}{r}"
                font, fill, al, border = cell.font, cell.fill, cell.alignment, cell.border
                entry: dict = {"v": _value(cell.value, cell.number_format)}
                if (sid := seg_of.get((ws.title, ref))) is not None:
                    entry["seg"] = sid
                if font is not None:
                    if font.b:
                        entry["b"] = True
                    if font.i:
                        entry["i"] = True
                    if (col := _color(font.color)) and col != "#000000":
                        entry["color"] = col
                if fill is not None and fill.fill_type == "solid" and (col := _color(fill.fgColor)):
                    entry["fill"] = col
                if al is not None:
                    if al.horizontal in ("center", "right", "centerContinuous"):
                        entry["align"] = "center" if al.horizontal.startswith("center") else "right"
                    if al.wrap_text:
                        entry["wrap"] = True
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and "align" not in entry:
                    entry["align"] = "right"
                if border is not None and any(getattr(border, s).style for s in ("left", "right", "top", "bottom") if getattr(border, s)):
                    entry["border"] = True
                row.append(entry)
            rows.append(row)
        widths = [round((ws.column_dimensions[get_column_letter(c)].width or 8.43) * 7 + 5) for c in range(1, cols_n + 1)]
        heights = {str(r): round(d.height * 4 / 3) for r, d in ws.row_dimensions.items() if d.height and r <= rows_n}
        merges = [str(m) for m in ws.merged_cells.ranges]
        objects = [{"kind": "picture", "anchor": _anchor(img)} for img in getattr(ws, "_images", [])]
        out.append({"name": ws.title, "rows": rows, "widths": widths, "heights": heights, "merges": merges, "objects": objects,
                    "truncated": ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS})
    charts = sum(1 for n in book.zip.namelist() if n.startswith("xl/charts/chart"))
    return {"sheets": out, "charts": charts}


def _anchor(img) -> str:
    try:
        m = img.anchor._from
        return f"{get_column_letter(m.col + 1)}{m.row + 1}"
    except Exception:  # noqa: BLE001
        return ""
