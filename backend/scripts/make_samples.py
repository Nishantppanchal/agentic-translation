"""Generate the Chinese and German test samples, their English references and the Chinese demo assets.

    cd backend && uv run python scripts/make_samples.py               # writes into ../samples
    cd backend && uv run python scripts/make_samples.py --out DIR     # writes the same tree into DIR

The Chinese samples continue the story of samples/yanshou-baogao.pdf (北区变电站 / North District
Substation, project UW-2291, 110 kV switchgear renovation); the German one continues the Umspannwerk
Nord samples. Every document plants traps for the agents: ambiguous utility terms that are NOT in the
termbase (两票三制, 五防, 消缺, 三措一案, 运维单位/运维方, 保电), terms used before they are defined,
repeated sentences, numbers/dates/amounts for QA, a paragraph across a page break, and translation-
memory repeats of yanshou-baogao.pdf. samples/SAMPLES.md (written by hand) says where each trap is
and what the app should do with it.

Written (relative to the output folder):
    fucha-baogao.pdf, fucha-baogao-scanned.pdf      复查报告: re-inspection of yanshou-baogao's defects
    jiance-baogao.docx                              检测报告: earthing and insulation test report
    anquan-peixun.pptx                              安全培训: safety training, four slides
    quexian-qingdan.xlsx                            缺陷清单: defect list and a summary sheet with formulas
    vertragsauszug.pdf, vertragsauszug-scanned.pdf  German construction-contract excerpt
    demo/zh/termbase.csv                            zh→en project termbase (same format as demo/termbase.csv)
    demo/zh/memory.tmx                              confirmed translations of yanshou-baogao.pdf
    demo/zh/jiance-baogao.vendor.en.xlf             "vendor" XLIFF 2.0 of jiance-baogao.docx, planted errors
    reference/<sample>.en.md                        English reference of each sample (and of yanshou-baogao.pdf)

The output is deterministic (fixed timestamps, no randomness): running it twice gives identical bytes.
Only the backend's own dependencies are used: PyMuPDF, python-docx, python-pptx, lxml, and xlsxwriter (a
dependency of python-pptx) for the workbook, because openpyxl 3.1 writes every string inline and the sample
needs a shared-string table.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import io
import re
import sys
import unicodedata
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import NamedTuple

import pymupdf
import xlsxwriter
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches
from pptx.util import Pt as PptxPt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # the backend package, imported in vendor_xliff()

ROOT = Path(__file__).resolve().parents[2]
STAMP = dt.datetime(2026, 5, 20, 9, 30)  # creation time written into every file (keeps the bytes stable)
ZIP_TIME = (2026, 5, 20, 9, 30, 0)
PDF_TIME = "D:20260520093000Z"
TMX_TIME = "20260520T093000Z"
LINK_URL = "https://example.com/uw-2291/jiance-fangan"  # the test plan the Word report links to

HEADER_ZH = "北区变电站改造工程 · UW-2291 · 内部资料"
HEADER_EN = "North District Substation Renovation Project · UW-2291 · Internal"


class Block(NamedTuple):
    """One block of a PDF sample: the source text and its English reference.

    kind: title | meta | h | p | caption | sign | table (src/en are lists of rows) | figure (src is the
    image name in the PDF, en describes it for the reference).
    """

    kind: str
    src: object
    en: object


def markdown_table(rows: list[list[str]]) -> str:
    """The convention of the TMX files: a table is one unit, written as a markdown table."""
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    return "\n".join(lines + ["| " + " | ".join(r) + " |" for r in rows[1:]])


# =========================================================================== #
# Content: yanshou-baogao.pdf (existing; the translation memory of the zh demo)
# =========================================================================== #
# Source strings exactly as they are in the PDF's text layer (line wraps removed).

YS_TITLE = ("设备验收报告——北区变电站", "Equipment Acceptance Report – North District Substation")
YS_META = ("日期：2026年3月14日 ｜ 项目编号：UW-2291", "Date: 14 March 2026 | Project No.: UW-2291")
YS_SCOPE_H = ("一、验收范围", "1. Scope of Acceptance")
YS_SCOPE = (
    "本次验收范围为110千伏开关设备的更新改造，包括二次系统。施工按照工程量清单第4.2.7项执行。承包方已完整移交全部试验报告。",
    ("The scope of this acceptance covers the renovation of the 110 kV switchgear, including the secondary systems. "
     "The works were carried out in accordance with item 4.2.7 of the bill of quantities. "
     "The Contractor has handed over all test reports in full."),
)
YS_DEFECTS_H = ("二、发现的缺陷", "2. Defects Identified")
YS_TABLE_ZH = [
    ["序号", "描述", "整改期限"],
    ["1", "地下室电缆桥架未接地", "2026年4月30日"],
    ["2", "出线间隔标识不完整", "2026年4月15日"],
    ["3", "UPS直流母线遗留工作", "2026年4月30日"],
]
YS_TABLE_EN = [
    ["No.", "Description", "Rectification deadline"],
    ["1", "Cable tray in the basement not earthed", "30 April 2026"],
    ["2", "Labelling of the outgoing feeder bays incomplete", "15 April 2026"],
    ["3", "Outstanding work on the UPS DC busbar", "30 April 2026"],
]
YS_FIGURE_EN = "Diagram: Distribution cabinet — UPS power supply — Cable tray; red label “Missing earthing connection”"
YS_CAPTION = ("图1：地下室设备概览", "Figure 1: Overview of the basement equipment")
YS_RESULT_H = ("三、结论", "3. Conclusion")
YS_RESULT = (
    "设备在上述缺陷整改完成前有条件验收。复查定于第18周进行。在此之前，保留结算金额5%的质量保证金。停送电操作须由电网调度中心统一下令。",
    ("The equipment is granted conditional acceptance pending rectification of the above defects. "
     "A re-inspection is scheduled for week 18. Until then, retention money of 5% of the settlement amount will be withheld. "
     "De-energisation and re-energisation may only be ordered centrally by the network control centre."),
)

YANSHOU = [
    Block("title", *YS_TITLE),
    Block("meta", *YS_META),
    Block("h", *YS_SCOPE_H),
    Block("p", *YS_SCOPE),
    Block("h", *YS_DEFECTS_H),
    Block("table", YS_TABLE_ZH, YS_TABLE_EN),
    Block("figure", "", YS_FIGURE_EN),
    Block("caption", *YS_CAPTION),
    Block("h", *YS_RESULT_H),
    Block("p", *YS_RESULT),
]

# =========================================================================== #
# Content: fucha-baogao.pdf (复查报告 — the re-inspection of yanshou-baogao's defects)
# =========================================================================== #

FUCHA_REPEAT = (  # the same paragraph twice (pages 1 and 2): one consistency group
    "乙方应在整改完成后5个工作日内书面通知甲方，并申请复验。",
    ("The Contractor shall notify the Client in writing within 5 working days of completing the rectification "
     "and request a re-inspection."),
)
FUCHA_SPLIT = (  # runs from the bottom of page 1 onto page 2
    ("第3项缺陷涉及UPS直流母线，相关工作须在UPS转旁路运行、相应馈线停电的条件下进行。根据运维单位报送的月度停电计划，"
     "2026年6月7日至6月10日为高考保电期，其间不安排计划停电。经甲乙双方协商，第3项缺陷的整改期限顺延至2026年6月15日。"
     "乙方应据此修订三措一案，经审批后于停电前3个工作日向电网调度中心提交停电申请。停电申请获批后，倒闸操作及现场安全措施"
     "由运维单位负责实施，乙方人员不得擅自操作任何一次设备。"),
    ("Defect No. 3 concerns the UPS DC busbar; the work can only be carried out with the UPS running on bypass and the "
     "associated feeders de-energised. According to the monthly outage schedule submitted by the O&M organisation, the "
     "period from 7 June to 10 June 2026 is a supply-assurance period for the national college entrance examination "
     "(gaokao), during which no planned outages are scheduled. Following consultation, the Client and the Contractor have "
     "agreed to extend the rectification deadline for defect No. 3 to 15 June 2026. The Contractor shall revise the Three "
     "Measures and One Plan accordingly and, once it has been approved, submit an outage request to the network control "
     "centre 3 working days before the outage. Once the outage request has been approved, the O&M organisation will carry "
     "out the switching operations and set up the safety measures on site; the Contractor's staff must not operate any "
     "primary equipment themselves."),
)
FUCHA_FIGURE_EN = (
    "Diagram (raster image, labels in Chinese): Distribution cabinet — UPS power supply — Cable tray, with an earth "
    "symbol labelled “Earthing bonding jumper”; green label “Earthing connection restored (measured 0.08 Ω)”"
)

FUCHA = [
    Block("title", "缺陷复查报告——北区变电站", "Defect Re-inspection Report – North District Substation"),
    Block("meta", "日期：2026年5月20日 ｜ 项目编号：UW-2291", "Date: 20 May 2026 | Project No.: UW-2291"),  # TM fuzzy
    Block("h", "一、复查范围", "1. Scope of Re-inspection"),
    Block(  # TM fuzzy: 验收 → 复查
        "p",
        "本次复查范围为110千伏开关设备的更新改造，包括二次系统。施工按照工程量清单第4.2.7项执行。承包方已完整移交全部试验报告。",
        "The scope of this re-inspection covers the renovation of the 110 kV switchgear, including the secondary systems. "
        "The works were carried out in accordance with item 4.2.7 of the bill of quantities. "
        "The Contractor has handed over all test reports in full.",
    ),
    Block(
        "p",
        "复查工作依据经审批的三措一案组织实施，现场工作时间为2026年5月20日9:30至14:30。甲方、乙方、监理单位及运维单位均派代表参加。",
        "The re-inspection was organised in accordance with the approved Three Measures and One Plan; work on site took "
        "place on 20 May 2026 from 9:30 to 14:30. The Client, the Contractor, the construction supervisor and the O&M "
        "organisation all sent representatives.",
    ),
    Block(
        "p",
        "复查组查阅了乙方提交的整改报告、接地导通测试记录及隐蔽工程影像资料，并对验收时发现的三项缺陷逐项进行了现场核查。"
        "对已整改部位，复查组按不低于30%的比例进行了抽样复测，复测结果与乙方记录一致。",
        "The re-inspection team reviewed the rectification report, the earthing continuity test records and the photographs "
        "of concealed works submitted by the Contractor, and checked each of the three defects found at acceptance on site. "
        "Where defects had been rectified, the team repeated at least 30% of the measurements on a sample basis; the results "
        "agreed with the Contractor's records.",
    ),
    Block("h", *YS_DEFECTS_H),  # TM exact
    Block("table", YS_TABLE_ZH, YS_TABLE_EN),  # TM exact
    Block(
        "p",
        "经现场核查，第1项缺陷已消缺：电缆桥架已加装16 mm²黄绿双色铜芯接地跨接线，实测接地导通电阻为0.08 Ω，满足不大于0.2 Ω的要求。"
        "第2项缺陷已整改完毕，出线间隔标识齐全、清晰。第3项缺陷尚未消缺。",
        "The on-site check found that defect No. 1 has been cleared: a 16 mm² green/yellow copper earthing bonding jumper has "
        "been fitted to the cable tray, and the measured earthing continuity resistance is 0.08 Ω, which meets the "
        "requirement of no more than "
        "0.2 Ω. Rectification of defect No. 2 is complete: the labels of the outgoing feeder bays are complete and legible. "
        "Defect No. 3 has not yet been cleared.",
    ),
    Block("p", *FUCHA_REPEAT),
    Block("figure", "fucha-fig1.png", FUCHA_FIGURE_EN),
    Block("caption", *YS_CAPTION),  # TM exact
    Block("p", *FUCHA_SPLIT),
    Block("h", *YS_RESULT_H),  # TM exact
    Block("p", "2026年3月14日设备验收报告的结论如下：", "The equipment acceptance report of 14 March 2026 concluded as follows:"),
    Block("p", *YS_RESULT),  # TM exact
    Block(
        "p",
        "本次复查确认第1、2项缺陷已闭环。经甲方同意，先行退还质保金¥12,500.00（结算金额的2%）；"
        "其余3%即¥18,750.00，待第3项缺陷消缺并经复查合格后退还。",
        "This re-inspection confirms that defects No. 1 and No. 2 have been closed out. With the Client's consent, retention "
        "money of CNY 12,500.00 (2% of the settlement amount) will be released now; the remaining 3%, i.e. CNY 18,750.00, "
        "will be released once defect No. 3 has been cleared and has passed re-inspection.",
    ),
    Block("p", *FUCHA_REPEAT),
    Block("h", "四、附注", "4. Notes"),
    Block(  # the forward definition: 三措一案 is used on page 1 and only defined here
        "p",
        "本报告所称“三措一案”，是指乙方针对停电作业编制的组织措施、技术措施、安全措施和施工方案，须经运维单位审核、甲方批准后方可实施。",
        "In this report, the “Three Measures and One Plan” means the organisational, technical and safety measures and the "
        "construction plan prepared by the Contractor for the outage work; they may only be implemented after review by "
        "the O&M organisation and approval by the Client.",
    ),
    Block("h", "五、签字确认", "5. Signatures"),
    Block("sign", "甲方代表：王建国", "Client representative: Wang Jianguo"),
    Block("sign", "乙方代表：李志强", "Contractor representative: Li Zhiqiang"),
    Block("sign", "运维单位代表：张伟", "O&M organisation representative: Zhang Wei"),
]

FUCHA_NOTES = """\
- **三措一案** (p. 1, twice) is only defined on p. 2 (“本报告所称…”): organisational, technical and safety
  measures plus the construction plan. Rendered “Three Measures and One Plan” throughout.
- **消缺** (defect elimination, utility jargon) is rendered “cleared”; **整改** (termbase) is “rectification”.
  The report uses both for the same defects: a translator may ask whether they differ.
- **运维单位** is the operation and maintenance organisation (here the utility's O&M department, not a
  contractor): “O&M organisation”.
- **保电 / 高考保电期**: a period of guaranteed supply (here during the national college entrance
  examination, gaokao) in which no planned outages are allowed: “supply-assurance period”.
- **监理单位** is the construction supervisor (the independent site supervision firm).
- Amounts: ¥ is written “CNY”. 2% of the settlement amount = CNY 12,500.00, so the settlement amount is
  CNY 625,000.00 and the 5% retention CNY 31,250.00 (the remaining 3% = CNY 18,750.00)."""

# =========================================================================== #
# Content: vertragsauszug.pdf (German contract excerpt)
# =========================================================================== #

VERTRAG_REPEAT = (  # § 4 (3) and § 5 (3): the same clause twice
    "(3) Die Anzeige bedarf der Schriftform und ist an die Projektleitung des AG zu richten.",
    "(3) The notice must be given in writing and addressed to the Client's project management.",
)
VERTRAG_SPLIT = (
    ("(1) Nach Fertigstellung der Leistung findet eine förmliche Abnahme statt; eine fiktive Abnahme nach § 12 Abs. 5 "
     "VOB/B ist ausgeschlossen. Der AN zeigt die Fertigstellung schriftlich an. Die Abnahme ist innerhalb von 12 Werktagen "
     "nach Zugang der Fertigstellungsanzeige beim AG durchzuführen. Über die Abnahme wird ein gemeinsames Protokoll gefertigt, in "
     "das vorbehaltene Mängel, vorbehaltene Vertragsstrafen sowie etwaige Einwendungen des AN aufzunehmen sind."),
    ("(1) Formal acceptance shall take place once the works have been completed; deemed acceptance under § 12(5) VOB/B "
     "is excluded. The Contractor shall give written notice of completion. Acceptance shall take place within 12 working "
     "days of the Client's receipt of the notice of completion. A joint record of acceptance shall be drawn up, stating any defects and "
     "contractual penalties that are reserved and any objections raised by the Contractor."),
)

VERTRAG = [
    Block("title", "Auszug aus dem Bauvertrag – Umspannwerk Nord", "Extract from the Construction Contract – Substation North"),
    Block("meta", "Vertrag Nr. UW-2291-BV vom 12.01.2026 | Projekt-Nr.: UW-2291",
          "Contract No. UW-2291-BV of 12 January 2026 | Project No.: UW-2291"),
    Block(
        "p",
        "Der folgende Auszug gibt die für die Bauausführung maßgeblichen Regelungen des Bauvertrags zwischen dem AG und dem "
        "AN über die Erneuerung der 110-kV-Schaltanlage im Umspannwerk Nord wieder. Ergänzend gilt die VOB/B in der bei "
        "Vertragsschluss gültigen Fassung.",
        "The following extract reproduces the provisions of the construction contract between the Client and the "
        "Contractor for the renewal of the 110 kV switchgear at Substation North that govern the execution of the works. "
        "The VOB/B, in the version applicable when the contract was concluded, applies in addition.",
    ),
    Block("h", "§ 4 Ausführungsfristen", "§ 4 Time for Completion"),
    Block(
        "p",
        "(1) Der AN beginnt mit den Arbeiten am 02.02.2026 und stellt die Leistung bis zum 13.03.2026 abnahmereif fertig. "
        "Beide Termine sind verbindliche Vertragsfristen im Sinne von § 5 VOB/B. Spätestens 10 Werktage vor Beginn legt der "
        "AN dem AG einen mit der Netzleitstelle abgestimmten Bauzeitenplan vor.",
        "(1) The Contractor shall commence the works on 2 February 2026 and complete them, ready for acceptance, by "
        "13 March 2026. Both dates are binding contractual deadlines within the meaning of § 5 VOB/B. No later than "
        "10 working days before the start, the Contractor shall submit to the Client a construction schedule agreed with "
        "the network control centre.",
    ),
    Block(
        "p",
        "(2) Sieht sich der AN in der ordnungsgemäßen Ausführung der Leistung behindert, so hat er dies dem AG unverzüglich "
        "anzuzeigen (Behinderungsanzeige nach § 6 Abs. 1 VOB/B). Unterlässt er die Anzeige, so hat er nur dann Anspruch auf "
        "Berücksichtigung der hindernden Umstände, wenn dem AG deren Tatsache und hindernde Wirkung offenkundig bekannt waren.",
        "(2) If the Contractor considers that it is being hindered in the proper execution of the works, it shall notify "
        "the Client without delay (notice of hindrance under § 6(1) VOB/B). If it fails to give notice, it is entitled to "
        "have the hindering circumstances taken into account only if the Client was clearly aware of the facts and of "
        "their hindering effect.",
    ),
    Block("p", *VERTRAG_REPEAT),
    Block("h", "§ 5 Bedenken gegen die Art der Ausführung", "§ 5 Concerns About the Method of Execution"),
    Block(
        "p",
        "(1) Hat der AN Bedenken gegen die vorgesehene Art der Ausführung, gegen die Güte der vom AG gelieferten Stoffe oder "
        "Bauteile oder gegen die Leistungen anderer Unternehmer, so hat er sie dem AG unverzüglich mitzuteilen "
        "(Bedenkenanzeige nach § 4 Abs. 3 VOB/B).",
        "(1) If the Contractor has concerns about the intended method of execution, about the quality of the materials or "
        "components supplied by the Client, or about the work of other contractors, it shall inform the Client without "
        "delay (notice of concerns under § 4(3) VOB/B).",
    ),
    Block(
        "p",
        "(2) Hält der AG trotz einer ordnungsgemäßen Bedenkenanzeige an seiner Anordnung fest, so haftet der AN nicht für "
        "die daraus entstehenden Mängel.",
        "(2) If the Client maintains its instruction despite a properly submitted notice of concerns, the Contractor is not "
        "liable for any defects resulting from it.",
    ),
    Block("p", *VERTRAG_REPEAT),
    Block("h", "§ 6 Vergütung und Zahlung", "§ 6 Remuneration and Payment"),
    Block(
        "p",
        "(1) Die Vergütung für die Leistungen nach LV-Position 4.2.7 beträgt pauschal 1.248.500,00 EUR netto zuzüglich der "
        "gesetzlichen Umsatzsteuer von derzeit 19 %.",
        "(1) The remuneration for the works under LV position 4.2.7 is a lump sum of EUR 1,248,500.00 net, plus statutory "
        "VAT, currently 19%.",
    ),
    Block(
        "p",
        "(2) Abschlagszahlungen sind binnen 21 Kalendertagen nach Zugang einer prüffähigen Abschlagsrechnung zu leisten. "
        "Die Schlusszahlung ist spätestens 30 Tage nach Zugang der prüffähigen Schlussrechnung fällig.",
        "(2) Interim payments shall be made within 21 calendar days of receipt of a verifiable interim invoice. The final "
        "payment is due no later than 30 days after receipt of the verifiable final invoice.",
    ),
    Block(
        "p",
        "(3) Überschreitet der AN die Frist nach § 4 Abs. 1 schuldhaft, so hat er für jeden Werktag des Verzugs eine "
        "Vertragsstrafe von 0,2 % der Nettoauftragssumme zu zahlen, insgesamt jedoch höchstens 5 % der Nettoauftragssumme.",
        "(3) If the Contractor culpably exceeds the deadline under § 4(1), it shall pay a contractual penalty of 0.2% of "
        "the net contract sum for each working day of delay, but no more than 5% of the net contract sum in total.",
    ),
    Block("h", "§ 7 Abnahme", "§ 7 Acceptance"),
    Block("p", *VERTRAG_SPLIT),
    Block(
        "p",
        "(2) Mit der Abnahme geht die Gefahr auf den AG über. Die Gewährleistungsfrist beginnt mit dem Tag der Abnahme und "
        "beträgt 5 Jahre.",
        "(2) Risk passes to the Client on acceptance. The warranty period begins on the day of acceptance and lasts 5 years.",
    ),
    Block("h", "§ 8 Sicherheitsleistung", "§ 8 Security"),
    Block(
        "p",
        "(1) Als Sicherheit für die Mängelansprüche behält der AG 5 % der Schlussrechnungssumme ein (Sicherheitseinbehalt). "
        "Der AN ist berechtigt, den Sicherheitseinbehalt durch eine selbstschuldnerische Bürgschaft eines in der "
        "Europäischen Union zugelassenen Kreditinstituts abzulösen.",
        "(1) As security for claims for defects, the Client shall retain 5% of the final invoice total (security "
        "retention). The Contractor is entitled to replace the security retention with an absolute guarantee (without the "
        "defence of prior recourse) from a credit institution authorised in the European Union.",
    ),
    Block(
        "p",
        "(2) Der nicht verwertete Sicherheitseinbehalt ist nach Ablauf der Gewährleistungsfrist zurückzugeben, spätestens "
        "jedoch am 31.03.2031.",
        "(2) Any part of the security retention that has not been drawn upon shall be returned when the warranty period "
        "expires, but no later than 31 March 2031.",
    ),
    Block("h", "Anlage 1 – Begriffe und Abkürzungen", "Annex 1 – Terms and Abbreviations"),
    Block("p", "Auftraggeber (AG): Acme Energie Netz GmbH, vertreten durch Dipl.-Ing. M. Weber.",
          "Client (AG): Acme Energie Netz GmbH, represented by Dipl.-Ing. M. Weber."),
    Block("p", "Auftragnehmer (AN): Albrecht Elektrotechnik GmbH, vertreten durch K. Albrecht.",
          "Contractor (AN): Albrecht Elektrotechnik GmbH, represented by K. Albrecht."),
    Block("p", "VOB/B: Vergabe- und Vertragsordnung für Bauleistungen, Teil B, Ausgabe 2016.",
          "VOB/B: German Construction Contract Procedures (Vergabe- und Vertragsordnung für Bauleistungen), Part B, "
          "2016 edition."),
    Block("p", "LV: Leistungsverzeichnis des AG, Stand 05.12.2025.",
          "LV: the Client's bill of quantities (Leistungsverzeichnis), as of 5 December 2025."),
]

VERTRAG_NOTES = """\
- **AG / AN** are used from the first paragraph but only defined in Annex 1 on p. 2: Auftraggeber (AG) =
  the Client, Auftragnehmer (AN) = the Contractor. Rendered “the Client” / “the Contractor” throughout
  (the abbreviations are kept only in Annex 1, where they are defined).
- **Abnahme** (termbase: acceptance) is a legal act under German construction law: *förmliche Abnahme* =
  formal acceptance, *fiktive Abnahme* = deemed acceptance; *abnahmereif* = ready for acceptance.
- **Behinderungsanzeige** (§ 6(1) VOB/B) = notice of hindrance (also “notice of obstruction/delay”).
- **Bedenkenanzeige** (§ 4(3) VOB/B) = notice of concerns (the contractor's written reservations about the
  specified method, materials or prior work).
- **Sicherheitseinbehalt** (termbase: security retention); *selbstschuldnerische Bürgschaft* = an absolute
  guarantee, enforceable without first pursuing the principal debtor.
- **VOB/B** and **LV** are kept as German abbreviations (as in the project style guide: codes unchanged).
- Amounts: 1.248.500,00 EUR → EUR 1,248,500.00 (decimal point in English)."""

# =========================================================================== #
# Content: jiance-baogao.docx (检测报告)
# =========================================================================== #
# A paragraph is a list of runs (text, format): "" plain, "b" bold, "i" italic, "link" hyperlink,
# "fn" footnote reference, "del"/"ins" tracked deletion/insertion.

JIANCE_CONDITIONS = ("测试条件：天气晴，环境温度25 ℃，相对湿度45%。",
                     "Test conditions: clear weather, ambient temperature 25 °C, relative humidity 45%.")
JIANCE_TABLE_HEAD = (["测量位置", "实测值", "标准要求", "结论"], ["Measuring point", "Measured value", "Requirement", "Result"])
JIANCE_FOOTNOTE = ("DL/T 475—2017《接地装置特性参数测量导则》。",
                   "DL/T 475-2017, *Guide for Measuring the Characteristic Parameters of Earthing Systems*.")

JIANCE = [
    ("h1", "接地及绝缘电阻检测报告", "Earthing and Insulation Resistance Test Report"),
    ("p", [("受甲方委托，乙方于2026年5月18日对北区变电站110 kV开关室的", ""), ("主接地网", "b"),
           ("进行了检测。实测接地电阻为0.42 Ω，", ""), ("满足设计值0.5 Ω的要求", "i"), ("。", "")],
     ("On behalf of the Client, the Contractor tested the **main earth grid** in the 110 kV switchgear room of the North "
      "District Substation on 18 May 2026. The measured earthing resistance was 0.42 Ω, *which meets the design value of "
      "0.5 Ω*.")),
    ("h2", "一、检测依据", "1. Basis of Testing"),
    ("p", [("检测按照经审批的", ""), ("检测方案", "link"), ("及DL/T 475—2017执行。", ""), ("", "fn")],
     f"The tests were carried out in accordance with the approved [test plan]({LINK_URL}) and DL/T 475-2017.[^1]"),
    ("p", [("检测前，运维单位已按操作票完成相关倒闸操作，并在被试设备两侧装设接地线。", "")],
     ("Before testing, the O&M organisation carried out the necessary switching operations in accordance with the "
      "switching order and applied earthing leads on both sides of the equipment under test.")),
    ("h2", "二、检测结果", "2. Test Results"),
    ("h3", "（一）接地电阻及接地导通测试", "2.1 Earthing resistance and earthing continuity"),
    ("p", [(JIANCE_CONDITIONS[0], "")], JIANCE_CONDITIONS[1]),
    ("table",
     [JIANCE_TABLE_HEAD[0], ["主接地网", "0.42 Ω", "≤0.5 Ω", "合格"], ["地下室电缆桥架", "0.08 Ω", "≤0.2 Ω", "合格"],
      ["UPS屏柜", "0.35 Ω", "≤0.2 Ω", "不合格"]],
     [JIANCE_TABLE_HEAD[1], ["Main earth grid", "0.42 Ω", "≤ 0.5 Ω", "Pass"], ["Basement cable tray", "0.08 Ω", "≤ 0.2 Ω", "Pass"],
      ["UPS cabinet", "0.35 Ω", "≤ 0.2 Ω", "Fail"]]),
    ("h3", "（二）绝缘电阻测试", "2.2 Insulation resistance"),
    ("p", [(JIANCE_CONDITIONS[0], "")], JIANCE_CONDITIONS[1]),
    ("table",
     [JIANCE_TABLE_HEAD[0], ["110 kV开关柜A相", "2500 MΩ", "≥1000 MΩ", "合格"], ["110 kV开关柜B相", "2300 MΩ", "≥1000 MΩ", "合格"],
      ["110 kV开关柜C相", "2400 MΩ", "≥1000 MΩ", "合格"]],
     [JIANCE_TABLE_HEAD[1], ["110 kV switchgear panel, phase A", "2500 MΩ", "≥ 1000 MΩ", "Pass"],
      ["110 kV switchgear panel, phase B", "2300 MΩ", "≥ 1000 MΩ", "Pass"],
      ["110 kV switchgear panel, phase C", "2400 MΩ", "≥ 1000 MΩ", "Pass"]]),
    ("h2", "三、结论", "3. Conclusion"),
    ("p", [("UPS屏柜接地导通电阻超标，判定为不合格。乙方须", ""), ("尽快", "del"), ("于2026年6月15日前", "ins"),
           ("完成消缺并复测。", "")],
     ("The earthing continuity resistance of the UPS cabinet exceeds the limit, so the cabinet is assessed as failed. "
      "The Contractor must clear the defect and repeat the test by 15 June 2026.")),
    ("p", [("其余检测项目均合格，相关设备具备投运条件。", "")],
     "All other test items passed; the equipment concerned is ready to be put into service."),
    ("p", [("检测人员：赵磊", "")], "Tested by: Zhao Lei"),
    ("p", [("审核：王建国", "")], "Reviewed by: Wang Jianguo"),
]

# The vendor's English, keyed by the segment source exactly as office.extract reads it. Planted errors are marked.
JIANCE_VENDOR = {
    "第 <x id=\"1\"/>1<x id=\"2\"/> 页": "Page <x id=\"1\"/>1<x id=\"2\"/>",
    HEADER_ZH: HEADER_EN,
    "接地及绝缘电阻检测报告": "Earthing and Insulation Resistance Test Report",
    # planted: forbidden term "Party A" (termbase: 甲方 → Client)
    "受甲方委托，乙方于2026年5月18日对北区变电站110 kV开关室的<g id=\"1\">主接地网</g>进行了检测。"
    "实测接地电阻为0.42 Ω，<g id=\"2\">满足设计值0.5 Ω的要求</g>。":
        "Commissioned by Party A, the Contractor tested the <g id=\"1\">main earth grid</g> in the 110 kV switchgear room of "
        "the North District Substation on 18 May 2026. The measured earthing resistance was 0.42 Ω, "
        "<g id=\"2\">meeting the design value of 0.5 Ω</g>.",
    "一、检测依据": "1. Basis of testing",
    # planted: the footnote mark <x id="2"/> is missing (the import refuses this segment)
    "检测按照经审批的<g id=\"1\">检测方案</g>及DL/T 475—2017执行。<x id=\"2\"/>":
        "Testing was carried out in accordance with the approved <g id=\"1\">test plan</g> and DL/T 475-2017.",
    "检测前，运维单位已按操作票完成相关倒闸操作，并在被试设备两侧装设接地线。":
        "Before testing, the O&M unit completed the relevant switching operations according to the switching order and "
        "installed earthing leads on both sides of the equipment under test.",
    "二、检测结果": "2. Test results",
    "（一）接地电阻及接地导通测试": "2.1 Earthing resistance and earthing continuity test",
    "测量位置": "Measuring point",
    "实测值": "Measured value",
    "标准要求": "Requirement",
    "结论": "Result",
    "主接地网": "Main earth grid",
    "0.42 Ω": "0.42 Ω",
    "≤0.5 Ω": "≤ 0.5 Ω",
    "合格": "Pass",
    "地下室电缆桥架": "Basement cable tray",
    "0.08 Ω": "0.08 Ω",
    "≤0.2 Ω": "≤ 0.2 Ω",
    "UPS屏柜": "UPS cabinet",
    "0.35 Ω": "0.53 Ω",  # planted: wrong number (transposed digits)
    "不合格": "Fail",
    "（二）绝缘电阻测试": "2.2 Insulation resistance test",
    "110 kV开关柜A相": "110 kV switchgear panel, phase A",
    "2500 MΩ": "2500 MΩ",
    "≥1000 MΩ": "≥ 1000 MΩ",
    "110 kV开关柜B相": "110 kV switchgear panel, phase B",
    "2300 MΩ": "2300 MΩ",
    "110 kV开关柜C相": "110 kV switchgear panel, phase C",
    "2400 MΩ": "2400 MΩ",
    "三、结论": "3. Conclusion",
    # planted: wrong date (15 June in the source; the date is the tracked insertion)
    "UPS屏柜接地导通电阻超标，判定为不合格。乙方须于2026年6月15日前完成消缺并复测。":
        "The earthing continuity resistance of the UPS cabinet exceeds the limit and is assessed as failed. The Contractor "
        "must eliminate the defect and retest by 15 May 2026.",
    # planted: left untranslated
    "其余检测项目均合格，相关设备具备投运条件。": "其余检测项目均合格，相关设备具备投运条件。",
    "检测人员：赵磊": "Tested by: Zhao Lei",
    "审核：王建国": "Checked by: Wang Jianguo",
    "<x id=\"1\"/>DL/T 475—2017《接地装置特性参数测量导则》。":
        "<x id=\"1\"/>DL/T 475-2017, Guide for Measuring the Characteristic Parameters of Earthing Systems.",
}
# planted: the repeated sentence translated two different ways (first and second occurrence)
JIANCE_VENDOR_CONDITIONS = [
    "Test conditions: clear weather, ambient temperature 25 °C, relative humidity 45%.",
    "Test conditions: sunny, ambient temp. 25 °C, RH 45%.",
]

JIANCE_NOTES = """\
- **Tracked changes:** the conclusion paragraph replaces 尽快 (“as soon as possible”, a tracked deletion) by
  于2026年6月15日前 (“by 15 June 2026”, a tracked insertion). The app reads tracked changes as accepted, so the
  source to translate contains the date and not 尽快.
- **运维单位** = the O&M organisation (operation and maintenance); **消缺** = clear / eliminate the defect.
- **操作票** = switching order (the written sequence for a switching operation); **倒闸操作** (termbase) =
  switching operation.
- **主接地网** = main earth grid; **接地导通** = earthing continuity (resistance between equipment and the grid).
- The sentence 测试条件：… appears twice (before each table) and must be translated identically.
- Footnote 1 (on DL/T 475—2017) is attached to the end of the paragraph in section 1."""

# =========================================================================== #
# Content: anquan-peixun.pptx (安全培训)
# =========================================================================== #

ANQUAN_INTERLOCK = ("严禁擅自解除防误闭锁装置。", "Never defeat an anti-maloperation interlock without authorisation.")

ANQUAN = [
    {
        "layout": 0,
        "title": ("停电作业安全培训", "Safety Training for Outage Work"),
        "body": [
            ("北区变电站110 kV开关设备改造工程（UW-2291）",
             "North District Substation 110 kV Switchgear Renovation Project (UW-2291)"),
            ("严格执行两票三制，全面落实五防要求",
             "Strict application of the Two Permits and Three Systems; full implementation of the Five Preventions"),
        ],
        "notes": [
            ("开场：本次培训面向参与UW-2291项目停电作业的全体人员，时长约45分钟。",
             "Opening: this training is for everyone involved in outage work on project UW-2291 and takes about 45 minutes."),
            ("培训结束后进行闭卷考试，80分及格。", "It ends with a closed-book test; the pass mark is 80 points."),
        ],
    },
    {
        "layout": 1,
        "title": ("作业前准备", "Before Work Starts"),
        "body": [
            ("开工前召开班前会，交代工作任务、安全措施和危险点",
             "Hold a pre-shift briefing before work starts: explain the tasks, the safety measures and the hazards"),
            ("倒闸操作由运维单位人员执行，施工人员不得擅自操作",
             "Switching operations are carried out by O&M staff only; construction staff must never operate equipment"),
            ("110 kV设备不停电时的安全距离为1.5 m", "Minimum safety clearance from live 110 kV equipment: 1.5 m"),
            ("6月7日至6月10日为保电期，暂停一切计划停电作业",
             "7 June to 10 June is a supply-assurance period: all planned outage work is suspended"),
        ],
        "notes": [
            ("保电期间如遇紧急缺陷，须经运维单位和电网调度中心批准后方可处理。",
             ("If an urgent defect arises during the supply-assurance period, it may only be dealt with after approval by "
              "the O&M organisation and the network control centre.")),
        ],
    },
    {
        "layout": 1,
        "title": ("什么是“五防”", "The “Five Preventions”"),
        "body": [
            ("防止误分、误合断路器", "Prevent opening or closing the wrong circuit breaker"),
            ("防止带负荷拉、合隔离开关", "Prevent opening or closing a disconnector under load"),
            ("防止带电挂（合）接地线（接地刀闸）",
             "Prevent applying an earthing lead (or closing an earthing switch) on live equipment"),
            ("防止带接地线（接地刀闸）合断路器（隔离开关）",
             "Prevent closing a circuit breaker (or disconnector) while an earthing lead (or earthing switch) is applied"),
            ("防止误入带电间隔", "Prevent entry into a live bay"),
        ],
        "notes": [
            ("“五防”是防止电气误操作的五项基本要求，由防误闭锁装置实现。",
             ("The Five Preventions are the five basic requirements for preventing electrical maloperation; they are "
              "enforced by anti-maloperation interlocks.")),
            ANQUAN_INTERLOCK,
        ],
    },
    {
        "layout": 1,
        "title": ("什么是“两票三制”", "The “Two Permits and Three Systems”"),
        "body": [
            ("两票：工作票、操作票", "Two permits: the work permit and the switching order"),
            ("三制：交接班制、巡回检查制、设备定期试验轮换制",
             "Three systems: shift handover, patrol inspection, and periodic testing and changeover of equipment"),
            ("无票作业一律按违章处理", "Any work without a permit is treated as a safety violation"),
        ],
        # A fixed-size text box (no autofit) that fits the Chinese exactly: the English needs about four times
        # as many characters, so the app must set it to shrink on overflow.
        "callout": ("安全第一、预防为主、综合治理", "Safety first, prevention foremost, comprehensive management"),
        "notes": [
            ("“两票三制”是电力安全生产的基本制度，本项目所有停电作业均须严格执行。",
             ("The Two Permits and Three Systems are the basic rules of safe operation in the power industry; all outage "
              "work on this project must follow them strictly.")),
            ANQUAN_INTERLOCK,
        ],
    },
]

ANQUAN_NOTES = """\
- **两票三制** and **五防** appear on slide 1 without explanation; slides 4 and 3 define them. Rendered
  “Two Permits and Three Systems” (work permit and switching order; shift handover, patrol inspection, and
  periodic testing and changeover of equipment) and “Five Preventions” (the five anti-maloperation rules).
- **保电** (slide 2) = a supply-assurance period in which no planned outages are allowed.
- **运维单位** = O&M organisation; **倒闸操作** (termbase) = switching operation; **操作票** = switching order.
- **安全第一、预防为主、综合治理** is the official safety policy slogan; the English is about four times as long
  as the Chinese, so the fixed-size text box on slide 4 overflows unless it shrinks.
- The note 严禁擅自解除防误闭锁装置。 appears on slides 3 and 4 and must be translated identically."""

# =========================================================================== #
# Content: quexian-qingdan.xlsx (缺陷清单)
# =========================================================================== #

QX_HEAD = [("编号", "ID"), ("位置", "Location"), ("缺陷描述", "Defect description"), ("责任方", "Responsible party"),
           ("整改期限", "Rectification deadline"), ("状态", "Status")]
QX_DONE, QX_OPEN = ("已整改", "Rectified"), ("未整改", "Not rectified")
QX_CONTRACTOR, QX_OM = ("乙方", "Contractor"), ("运维方", "O&M party")
QX_BASEMENT, QX_SWITCHROOM, QX_UPSROOM = ("地下室", "Basement"), ("110 kV开关室", "110 kV switchgear room"), ("UPS室", "UPS room")
QX_IMPORTANT = ("【重要】", "[Important]")  # the bold part of the rich-text cell
QX_ROWS = [  # (no., location, description, responsible, deadline, status); description None = the rich-text cell
    (1, QX_BASEMENT, ("地下室电缆桥架未接地", "Cable tray in the basement not earthed"), QX_CONTRACTOR, dt.date(2026, 4, 30), QX_DONE),
    (2, QX_SWITCHROOM, ("出线间隔标识不完整", "Labelling of the outgoing feeder bays incomplete"), QX_CONTRACTOR,
     dt.date(2026, 4, 15), QX_DONE),
    (3, QX_UPSROOM, ("UPS直流母线遗留工作", "Outstanding work on the UPS DC busbar"), QX_CONTRACTOR, dt.date(2026, 6, 15), QX_OPEN),
    (4, QX_UPSROOM, ("UPS屏柜接地导通电阻超标（实测0.35 Ω，要求≤0.2 Ω）",
                     "Earthing continuity resistance of the UPS cabinet above the limit (measured 0.35 Ω, required ≤ 0.2 Ω)"),
     QX_CONTRACTOR, dt.date(2026, 6, 15), QX_OPEN),
    (5, ("主控室", "Main control room"), ("后台监控系统频繁误报信号，需运维方配合消缺",
                                        ("Frequent false alarms on the station monitoring system; clearing the defect "
                                         "needs the O&M party's support")),
     QX_OM, dt.date(2026, 5, 31), QX_DONE),
    (6, QX_SWITCHROOM, None, QX_OM, dt.date(2026, 6, 30), QX_OPEN),
]
QX_RICH = ("2号主变110 kV侧隔离开关触头发热，红外测温82 ℃",
           "Contacts of the 110 kV-side disconnector of main transformer No. 2 overheating; infrared reading 82 °C")
QX_DATE_FORMAT = 'yyyy"年"m"月"d"日"'
QX_SUMMARY = [  # ((label, English), value or formula, number format, the value as the reference shows it)
    (("统计项目", "Item"), ("数值", "Value"), None, None),
    (("缺陷总数", "Total number of defects"), "=COUNTA(缺陷!A2:A7)", "0", "6"),
    (("已整改数量", "Defects with rectification completed"), '=COUNTIF(缺陷!F:F,"已整改")', "0", "3"),
    (("未整改数量", "Defects with rectification outstanding"), '=COUNTIF(缺陷!F:F,"未整改")', "0", "3"),
    (("整改完成率", "Rectification completion rate"), "=B3/B2", "0%", "50%"),
    (("乙方责任缺陷", "Defects for which the Contractor is responsible"), '=COUNTIF(缺陷!D:D,"乙方")', "0", "4"),
    (("结算金额（元）", "Settlement amount (CNY)"), 625000, "#,##0.00", "625,000.00"),
    (("质量保证金比例", "Retention money (share of the settlement amount)"), 0.05, "0%", "5%"),
    (("质量保证金（元）", "Retention money (CNY)"), "=B7*B8", "#,##0.00", "31,250.00"),
    (("统计日期", "Date of this summary"), dt.date(2026, 5, 20), QX_DATE_FORMAT, "20 May 2026"),
]

QUEXIAN_NOTES = """\
- **已整改 / 未整改** (status) are shared strings used three times each, and the summary sheet's COUNTIF
  formulas compare against them (as they compare against 乙方): translating those cells changes the
  formulas' results unless the formulas are updated too.
- **运维方** (responsible party) = the O&M party: the utility's operation and maintenance side, as opposed
  to 乙方, the Contractor. **消缺** = clear / eliminate a defect.
- The deadlines and the summary date are real Excel dates (displayed as 2026年4月30日); the settlement amount
  and the 5% rate are numbers. Row 6's description is rich text: 【重要】 is bold."""


# =========================================================================== #
# Helpers
# =========================================================================== #


def fixed_zip(data: bytes) -> bytes:
    """Rewrite an Office zip with fixed timestamps (entries and docProps/core.xml) so the bytes are stable."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2026-05-20T09:30:00Z", body)
            entry = zipfile.ZipInfo(info.filename, ZIP_TIME)
            entry.external_attr = 0o644 << 16
            z.writestr(entry, body, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def esc(text: str) -> str:
    return html.escape(text, quote=False)


def norm(text: str) -> str:
    """Text without whitespace, ligatures (ﬀ → ff) and width variants: for comparing against a text layer."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


# =========================================================================== #
# PDF
# =========================================================================== #

PAGE = pymupdf.paper_rect("a4")
BODY = pymupdf.Rect(56, 66, 539, 780)
PDF_CSS = """
* {font-family: sans-serif; font-size: 10.5pt; line-height: 1.55;}
h1 {font-size: 16pt; margin: 0 0 4pt 0;}
h2 {font-size: 12.5pt; margin: 12pt 0 4pt 0;}
p {margin: 0 0 6pt 0;}
p.meta {color: #333; margin-bottom: 8pt;}
p.caption {font-size: 9.5pt; text-align: center; margin-bottom: 10pt;}
p.figure {text-align: center; margin: 4pt 0 2pt 0;}
p.sign {margin: 0 0 2pt 0;}
table {border-collapse: collapse; width: 100%; margin: 2pt 0 8pt 0;}
td, th {border: 0.6pt solid #888; padding: 3pt 5pt; text-align: left; vertical-align: top;}
th {font-weight: normal; border-bottom: 1.4pt solid #444;}  /* no fill: Story repaints th fills on the next page */
"""
FURNITURE_CSS = "* {font-family: sans-serif; font-size: 8pt; color: #555;}"


def blocks_html(blocks: list[Block], figure_width: int) -> str:
    out = []
    for b in blocks:
        if b.kind == "table":
            head, *rows = b.src
            out.append("<table><tr>" + "".join(f"<th>{esc(c)}</th>" for c in head) + "</tr>"
                       + "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in rows) + "</table>")
        elif b.kind == "figure":
            out.append(f'<p class="figure"><img src="{b.src}" style="width: {figure_width}px"/></p>')
        elif b.kind == "title":
            out.append(f"<h1>{esc(b.src)}</h1>")
        elif b.kind == "h":
            out.append(f"<h2>{esc(b.src)}</h2>")
        else:
            cls = f' class="{b.kind}"' if b.kind in ("meta", "caption", "sign") else ""
            out.append(f"<p{cls}>{esc(b.src)}</p>")
    return "\n".join(out)


def flow_pdf(blocks: list[Block], *, title: str, header: str, footer: str, images: dict[str, bytes] | None = None,
             figure_width: int = 300) -> pymupdf.Document:
    """Typeset the blocks as flowing text over as many A4 pages as needed (a paragraph may run on to the next
    page), then add the running header and the page-number footer (`footer` takes {page} and {pages})."""
    archive = pymupdf.Archive()
    for name, data in (images or {}).items():
        archive.add(data, name)
    story = pymupdf.Story(blocks_html(blocks, figure_width), user_css=PDF_CSS, archive=archive)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    more = True
    while more:
        device = writer.begin_page(PAGE)
        more, _ = story.place(BODY)
        story.draw(device)
        writer.end_page()
    writer.close()
    doc = pymupdf.open("pdf", buf.getvalue())
    grey = (0.7, 0.7, 0.7)
    for number, page in enumerate(doc, start=1):
        page.insert_htmlbox(pymupdf.Rect(56, 30, 539, 46), f"<p>{esc(header)}</p>", css=FURNITURE_CSS)
        page.draw_line((56, 46), (539, 46), color=grey, width=0.5)
        page.draw_line((56, 798), (539, 798), color=grey, width=0.5)
        text = esc(footer.format(page=number, pages=len(doc)))
        page.insert_htmlbox(pymupdf.Rect(56, 802, 539, 818), f'<p style="text-align: right">{text}</p>', css=FURNITURE_CSS)
    doc.set_metadata({"title": title, "author": "UW-2291", "creator": "make_samples.py", "producer": "PyMuPDF",
                      "creationDate": PDF_TIME, "modDate": PDF_TIME})
    doc.subset_fonts()
    return doc


def pdf_bytes(doc: pymupdf.Document) -> bytes:
    return doc.tobytes(garbage=4, deflate=True, no_new_id=True)


def scanned(doc: pymupdf.Document) -> bytes:
    """Image-only copy (no text layer), rasterised like the existing -scanned.pdf samples."""
    scan = pymupdf.open()
    for page in doc:
        sp = scan.new_page(width=page.rect.width, height=page.rect.height)
        sp.insert_image(sp.rect, stream=page.get_pixmap(dpi=120).tobytes("jpeg", jpg_quality=80))
    scan.set_metadata({"title": doc.metadata["title"], "creator": "make_samples.py", "producer": "PyMuPDF",
                       "creationDate": PDF_TIME, "modDate": PDF_TIME})
    return scan.tobytes(garbage=4, deflate=True, no_new_id=True)


def check_runs_across(doc: pymupdf.Document, paragraph: str) -> None:
    """The paragraph must start on page 1 and end on page 2, broken mid-sentence (not after 。 or .)."""
    first, second, para = norm(doc[0].get_text()), norm(doc[1].get_text()), norm(paragraph)
    head = next((para[:n] for n in range(len(para), 0, -1) if first.find(para[:n]) >= 0), "")
    assert 10 < len(head) < len(para) and para[len(head):][:10] in second, "the paragraph doesn't run across the page break"
    assert head[-1] not in "。.;；:：", f"the page break falls between sentences: …{head[-12:]}"


def fucha_diagram() -> bytes:
    """The basement figure of yanshou-baogao.pdf after the re-inspection: a raster image with Chinese labels."""
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=230)
    label_css = "* {font-family: sans-serif; font-size: 14pt; text-align: center;}"
    for rect, label in ((pymupdf.Rect(20, 40, 140, 100), "配电柜"), (pymupdf.Rect(280, 40, 400, 100), "UPS电源"),
                        (pymupdf.Rect(140, 150, 260, 210), "电缆桥架")):
        page.draw_rect(rect, color=(0.1, 0.3, 0.7), fill=(0.9, 0.94, 1), width=2)
        page.insert_htmlbox(rect + (0, 19, 0, 0), f"<p>{label}</p>", css=label_css)
    for a, b in (((140, 70), (280, 70)), ((80, 100), (170, 150)), ((340, 100), (230, 150))):
        page.draw_line(a, b, color=(0.2, 0.2, 0.2), width=2)
    green = (0.05, 0.5, 0.15)
    page.draw_polyline([(260, 180), (305, 180), (305, 196)], color=green, width=2)  # bonding jumper to an earth symbol
    for half, y in ((12, 196), (8, 201), (4, 206)):
        page.draw_line((305 - half, y), (305 + half, y), color=green, width=2)
    small = "* {font-family: sans-serif; font-size: 12pt; color: #0d8026;}"
    page.insert_htmlbox(pymupdf.Rect(322, 188, 418, 206), "<p>接地跨接线</p>", css=small)
    page.insert_htmlbox(pymupdf.Rect(60, 8, 360, 28), "<p style='text-align: center'>接地连接已恢复（实测0.08 Ω）</p>",
                        css=small.replace("12pt", "13pt"))
    return page.get_pixmap(dpi=150).tobytes("png")


def build_fucha() -> pymupdf.Document:
    doc = flow_pdf(FUCHA, title="缺陷复查报告——北区变电站", header=HEADER_ZH, footer="第 {page} 页 共 {pages} 页",
                   images={"fucha-fig1.png": fucha_diagram()}, figure_width=370)
    assert len(doc) == 2, f"fucha-baogao.pdf has {len(doc)} pages"
    check_runs_across(doc, FUCHA_SPLIT[0])
    return doc


def build_vertrag() -> pymupdf.Document:
    doc = flow_pdf(VERTRAG, title="Auszug aus dem Bauvertrag – Umspannwerk Nord",
                   header="Umspannwerk Nord · Projekt UW-2291 · Vertragsauszug · Vertraulich", footer="Seite {page} von {pages}")
    assert len(doc) == 2, f"vertragsauszug.pdf has {len(doc)} pages"
    check_runs_across(doc, VERTRAG_SPLIT[0])
    return doc


# =========================================================================== #
# Word
# =========================================================================== #

FOOTNOTES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
<w:footnote w:id="1"><w:p><w:pPr><w:spacing w:after="0"/></w:pPr><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>\
<w:footnoteRef/></w:r><w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:t>{text}</w:t></w:r></w:p></w:footnote>
</w:footnotes>"""
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"


def w_el(tag: str, text: str | None = None, **attrs: str):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(key.replace("_", ":", 1)), value)
    if text is not None:
        el.text = text
        el.set(qn("xml:space"), "preserve")
    return el


def w_run(*children, rpr: list | None = None):
    run = w_el("w:r")
    if rpr:
        props = w_el("w:rPr")
        for child in rpr:
            props.append(child)
        run.append(props)
    for child in children:
        run.append(child)
    return run


def add_runs(paragraph, runs: list[tuple[str, str]]) -> None:
    for text, fmt in runs:
        if fmt == "link":
            rid = paragraph.part.relate_to(LINK_URL, RT.HYPERLINK, is_external=True)
            link = w_el("w:hyperlink", r_id=rid)
            link.append(w_run(w_el("w:t", text), rpr=[w_el("w:color", w_val="0563C1"), w_el("w:u", w_val="single")]))
            paragraph._p.append(link)
        elif fmt == "fn":
            paragraph._p.append(w_run(w_el("w:footnoteReference", w_id="1"), rpr=[w_el("w:vertAlign", w_val="superscript")]))
        elif fmt in ("del", "ins"):
            change = w_el(f"w:{fmt}", w_id="91" if fmt == "del" else "92", w_author="王建国", w_date="2026-05-19T10:00:00Z")
            change.append(w_run(w_el("w:delText" if fmt == "del" else "w:t", text)))
            paragraph._p.append(change)
        else:
            run = paragraph.add_run(text)
            if fmt == "b":
                run.bold = True
            elif fmt == "i":
                run.italic = True


def page_field(paragraph, before: str, after: str) -> None:
    """`before` PAGE `after` — a live page-number field (the result "1" is what Word shows until it updates)."""
    paragraph.add_run(before)
    for child in (w_el("w:fldChar", w_fldCharType="begin"), w_el("w:instrText", " PAGE "),
                  w_el("w:fldChar", w_fldCharType="separate"), w_el("w:t", "1"), w_el("w:fldChar", w_fldCharType="end")):
        paragraph._p.append(w_run(child))
    paragraph.add_run(after)


def chinese_styles(doc) -> None:
    """Chinese fonts (SimSun body, SimHei headings) and zh-CN as the document language."""
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = defaults.find(qn("w:lang"))
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    for name, font in (("Normal", "SimSun"), ("Heading 1", "SimHei"), ("Heading 2", "SimHei"), ("Heading 3", "SimHei")):
        rpr = doc.styles[name].element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = w_el("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.attrib.pop(qn("w:eastAsiaTheme"), None)  # a theme font would win over the explicit one
        rfonts.set(qn("w:eastAsia"), font)
    doc.styles["Heading 1"].font.size = Pt(16)


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    props.append(w_el("w:shd", w_val="clear", w_color="auto", w_fill=fill))


def build_jiance() -> bytes:
    doc = Document()
    chinese_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.header.paragraphs[0].text = HEADER_ZH
    page_field(section.footer.paragraphs[0], "第 ", " 页")
    for kind, src, _ in JIANCE:
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(src, level=int(kind[1]))
        elif kind == "p":
            add_runs(doc.add_paragraph(), src)
        elif kind == "table":
            table = doc.add_table(rows=len(src), cols=len(src[0]))
            table.style = "Table Grid"
            for r, row in enumerate(src):
                for c, text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = text
                    if r == 0:
                        cell.paragraphs[0].runs[0].bold = True
                        shade(cell, "E9E9E9")
            doc.add_paragraph()  # space after the table (empty: not a segment)
    footnotes = Part(PackURI("/word/footnotes.xml"), FOOTNOTES_CT, FOOTNOTES_XML.format(text=JIANCE_FOOTNOTE[0]).encode(),
                     doc.part.package)
    doc.part.relate_to(footnotes, RT.FOOTNOTES)
    props = doc.core_properties
    props.title, props.author, props.last_modified_by = "接地及绝缘电阻检测报告", "赵磊", "王建国"
    props.created = props.modified = STAMP
    props.revision = 2
    buf = io.BytesIO()
    doc.save(buf)
    return fixed_zip(buf.getvalue())


# =========================================================================== #
# PowerPoint
# =========================================================================== #


def build_anquan() -> bytes:
    prs = Presentation()  # the default 4:3 template, like unterweisung.pptx
    for spec in ANQUAN:
        slide = prs.slides.add_slide(prs.slide_layouts[spec["layout"]])
        slide.shapes.title.text = spec["title"][0]
        body = slide.placeholders[1].text_frame
        body.text = spec["body"][0][0]
        for zh, _ in spec["body"][1:]:
            body.add_paragraph().text = zh
        if callout := spec.get("callout"):
            body_ph = slide.placeholders[1]  # make room below the bullets (set all four: it inherits them otherwise)
            body_ph.left, body_ph.top, body_ph.width, body_ph.height = body_ph.left, body_ph.top, body_ph.width, Inches(3.6)
            box = slide.shapes.add_textbox(Inches(2.7), Inches(5.75), Inches(4.6), Inches(0.62))
            box.fill.solid()
            box.fill.fore_color.rgb = PptxColor(0xFF, 0xF2, 0xCC)
            box.line.color.rgb = PptxColor(0xC0, 0x00, 0x00)
            frame = box.text_frame
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.NONE  # fixed size: the English must shrink to fit
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = callout[0]
            run.font.size, run.font.bold = PptxPt(20), True
            run.font.color.rgb = PptxColor(0xC0, 0x00, 0x00)
        slide.notes_slide.notes_text_frame.text = "\n".join(zh for zh, _ in spec["notes"])
    props = prs.core_properties
    props.title, props.author, props.last_modified_by = "停电作业安全培训", "张伟", "张伟"
    props.created = props.modified = STAMP
    props.revision = 1
    buf = io.BytesIO()
    prs.save(buf)
    return fixed_zip(buf.getvalue())


# =========================================================================== #
# Excel
# =========================================================================== #


def build_quexian() -> bytes:
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    wb.set_properties({"title": "缺陷清单", "author": "李志强", "created": STAMP})
    base = {"border": 1, "border_color": "#999999", "valign": "top"}
    head = wb.add_format({**base, "bold": True, "bg_color": "#DDEBF7"})
    cell = wb.add_format(base)
    wrap = wb.add_format({**base, "text_wrap": True})
    date = wb.add_format({**base, "num_format": QX_DATE_FORMAT})
    bold = wb.add_format({"bold": True})

    ws = wb.add_worksheet("缺陷")
    ws.write_row(0, 0, [zh for zh, _ in QX_HEAD], head)
    for r, (no, (place, _), desc, (party, _), due, (status, _)) in enumerate(QX_ROWS, start=1):
        ws.write_number(r, 0, no, cell)
        ws.write_string(r, 1, place, cell)
        if desc:
            ws.write_string(r, 2, desc[0], wrap)
        else:  # rich text: the 【重要】 prefix is bold
            ws.write_rich_string(r, 2, bold, QX_IMPORTANT[0], QX_RICH[0], wrap)
        ws.write_string(r, 3, party, cell)
        ws.write_datetime(r, 4, dt.datetime.combine(due, dt.time()), date)
        ws.write_string(r, 5, status, cell)  # 已整改 / 未整改: shared strings, three uses each
    for col, width in enumerate((7, 15, 52, 10, 17, 10)):
        ws.set_column(col, col, width)
    ws.freeze_panes(1, 0)

    summary = wb.add_worksheet("汇总")
    for r, ((label, _), value, fmt, shown) in enumerate(QX_SUMMARY):
        style = head if r == 0 else wb.add_format({**base, "num_format": fmt})
        summary.write_string(r, 0, label, head if r == 0 else cell)
        if isinstance(value, tuple):
            summary.write_string(r, 1, value[0], style)
        elif isinstance(value, str):  # a formula, with its result cached for readers that don't recalculate
            cached = float(shown.rstrip("%").replace(",", "")) / (100 if shown.endswith("%") else 1)
            summary.write_formula(r, 1, value, style, cached)
        elif isinstance(value, dt.date):
            summary.write_datetime(r, 1, dt.datetime.combine(value, dt.time()), style)
        else:
            summary.write_number(r, 1, value, style)
    summary.set_column(0, 0, 20)
    summary.set_column(1, 1, 18)
    wb.close()
    return fixed_zip(buf.getvalue())


# =========================================================================== #
# Chinese demo assets: termbase, translation memory, vendor XLIFF
# =========================================================================== #

TERMBASE_ZH = [  # (Chinese, English, note, forbidden) — the ambiguous question terms are deliberately absent
    ("北区变电站", "North District Substation", "Site name — always translated", "North Area Substation|Beiqu Substation"),
    ("甲方", "Client", "Contract party — write “the Client”", "Party A"),
    ("乙方", "Contractor", "Contract party — write “the Contractor”", "Party B"),
    ("承包方", "Contractor", "The same party as 乙方", "contracting party|undertaker"),
    ("电缆桥架", "cable tray", "Metal tray carrying cables", "cable bridge|cable rack"),
    ("整改", "rectification", "Correcting a recorded defect (已整改 = rectified)", "reform|correction"),
    ("验收", "acceptance", "Formal acceptance of the works", "check and accept|check before acceptance"),
    ("接地", "earth", "British usage: earth / earthing / earthed — never ground", "grounding|grounded"),
    ("质量保证金", "retention money", "Share of the settlement amount withheld until the defects are cleared",
     "quality deposit|quality guarantee deposit|quality assurance deposit"),
    ("质保金", "retention money", "Short form of 质量保证金", "quality deposit|warranty deposit"),
    ("倒闸操作", "switching operation", "Opening and closing breakers and disconnectors to change the system state",
     "gate switching|switchover operation"),
    ("工程量清单", "bill of quantities", "", "engineering quantity list|quantity list"),
    ("出线间隔", "outgoing feeder bay", "", "outlet interval|outgoing line interval"),
    ("电网调度中心", "network control centre", "British spelling", "grid dispatching center|dispatch center"),
    ("绝缘电阻", "insulation resistance", "", "isolation resistance"),
]


def termbase_csv() -> bytes:
    rows = ["Chinese;English;Note;Forbidden"] + [";".join(row) for row in TERMBASE_ZH]
    for row in TERMBASE_ZH:
        assert not any(c in "".join(row) for c in ';,"'), f"a termbase cell contains a delimiter: {row}"
    return ("\n".join(rows) + "\n").encode()


def tm_units() -> list[tuple[str, str]]:
    """Confirmed translations of yanshou-baogao.pdf, one unit per paragraph / heading / table (like demo/memory.tmx)."""
    units = []
    for b in YANSHOU:
        if b.kind == "table":
            units.append((markdown_table(b.src), markdown_table(b.en)))
        elif b.kind != "figure":
            units.append((b.src, b.en))
    return units


def memory_tmx() -> bytes:
    xml_lang = "{http://www.w3.org/XML/1998/namespace}lang"
    root = etree.Element("tmx", version="1.4")
    etree.SubElement(root, "header", creationtool="AgenticTranslator", creationtoolversion="1", segtype="paragraph",
                     attrib={"o-tmf": "agentic-translator", "adminlang": "en", "srclang": "zh-Hans", "datatype": "plaintext"})
    body = etree.SubElement(root, "body")
    for source, target in tm_units():
        tu = etree.SubElement(body, "tu", creationdate=TMX_TIME, changedate=TMX_TIME)
        etree.SubElement(tu, "prop", type="x-origin").text = "yanshou-baogao.pdf (confirmed)"
        for lang, text in (("zh-Hans", source), ("en", target)):
            etree.SubElement(etree.SubElement(tu, "tuv", attrib={xml_lang: lang}), "seg").text = text
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", pretty_print=True)


def vendor_xliff(docx: bytes) -> bytes:
    """jiance-baogao.docx as a vendor returned it: XLIFF 2.0 with segment ids and codes from office.extract."""
    # read-only use of the app's filters, like make_vendor_xliff.py
    from translator.filters import office, xliff

    segments, _ = office.extract(docx, "docx")
    conditions = iter(JIANCE_VENDOR_CONDITIONS)
    for s in segments:
        if s["source"] == JIANCE_CONDITIONS[0]:
            s["target"] = next(conditions)
        elif s["translatable"]:
            s["target"] = JIANCE_VENDOR[s["source"]]
        s["status"] = "draft"
    return xliff.export(segments, src="zh-Hans", tgt="en", original="jiance-baogao.docx")


# =========================================================================== #
# English references
# =========================================================================== #

REF_INTRO = ("A professional English reference translation, for judging the app's output without reading the source. "
             "Conventions: British spelling, dates as “20 May 2026”, decimal point, the project termbase applied. "
             "Generated by `backend/scripts/make_samples.py`; `samples/SAMPLES.md` lists the traps this sample plants.")


def reference(name: str, about: str, body: str, notes: str) -> str:
    return f"# English reference: {name}\n\n{about}\n\n{REF_INTRO}\n\n---\n\n{body.strip()}\n\n---\n\n## Translator's notes\n\n{notes}\n"


def markdown_of_blocks(blocks: list[Block]) -> str:
    out = []
    for b in blocks:
        if b.kind == "title":
            out.append(f"# {b.en}")
        elif b.kind == "h":
            out.append(f"## {b.en}")
        elif b.kind == "table":
            out.append(markdown_table(b.en))
        elif b.kind == "figure":
            out.append(f"*[{b.en}]*")
        elif b.kind == "sign":
            out.append(f"{b.en}  ")
        else:
            out.append(b.en)
    return "\n\n".join(out).replace("  \n\n", "  \n")


def ref_pdf(blocks: list[Block], header: str = "", footer: str = "") -> str:
    furniture = f"> Running header (every page): {header}  \n> Footer: {footer}\n\n" if header else ""
    return furniture + markdown_of_blocks(blocks)


def ref_jiance() -> str:
    out = [f"> Header: {HEADER_EN}  \n> Footer: Page *n* (a live PAGE field)"]
    for kind, _, en in JIANCE:
        if kind == "table":
            out.append(markdown_table(en))
        elif kind in ("h1", "h2", "h3"):
            out.append("#" * int(kind[1]) + " " + en)
        else:
            out.append(en)
    out.append(f"[^1]: {JIANCE_FOOTNOTE[1]}")
    return "\n\n".join(out)


def ref_anquan() -> str:
    out = []
    for number, spec in enumerate(ANQUAN, start=1):
        out.append(f"## Slide {number}: {spec['title'][1]}")
        if spec["layout"] == 0:  # title slide: subtitle lines
            out.append("  \n".join(en for _, en in spec["body"]))
        else:
            out.append("\n".join(f"- {en}" for _, en in spec["body"]))
        if callout := spec.get("callout"):
            out.append(f"Text box (fixed size, bold red on yellow): **{callout[1]}**")
        out.append("Speaker notes:\n\n" + "\n>\n".join(f"> {en}" for _, en in spec["notes"]))
    return "\n\n".join(out)


def ref_quexian() -> str:
    rows = [[en for _, en in QX_HEAD]]
    for no, place, desc, party, due, status in QX_ROWS:
        text = desc[1] if desc else f"**{QX_IMPORTANT[1]}** {QX_RICH[1]}"
        rows.append([str(no), place[1], text, party[1], f"{due.day} {due:%B %Y}", status[1]])
    summary = [["Item", "Value", "Formula"]] + [
        [label[1], shown, f"`{value}`" if isinstance(value, str) and value.startswith("=") else ""]
        for (label, value, _, shown) in QX_SUMMARY[1:]
    ]
    return ("## Sheet “缺陷” (Defects)\n\n" + markdown_table(rows)
            + "\n\n## Sheet “汇总” (Summary)\n\n" + markdown_table(summary)
            + "\n\nThe formulas stay as they are; translated sheet names and status texts would break them "
              "(`缺陷!F:F` compares against “已整改” / “未整改”).")


def references() -> dict[str, str]:
    pdf_about = "Source: `samples/{0}.pdf` ({1}, {2}; `samples/{0}-scanned.pdf` is the same document as an image-only scan)."
    yanshou_notes = ("- These translations are exactly the units of `samples/demo/zh/memory.tmx`.\n"
                     "- **质量保证金** (termbase) = retention money; **电网调度中心** (termbase) = network control centre.")
    return {
        "yanshou-baogao.en.md": reference(
            "yanshou-baogao.pdf", pdf_about.format("yanshou-baogao", "Chinese", "1 page"),
            ref_pdf(YANSHOU), yanshou_notes),
        "fucha-baogao.en.md": reference(
            "fucha-baogao.pdf", pdf_about.format("fucha-baogao", "Chinese", "2 pages"),
            ref_pdf(FUCHA, HEADER_EN, "Page 1 of 2 / Page 2 of 2"), FUCHA_NOTES),
        "vertragsauszug.en.md": reference(
            "vertragsauszug.pdf", pdf_about.format("vertragsauszug", "German", "2 pages"),
            ref_pdf(VERTRAG, "Substation North · Project UW-2291 · Contract Extract · Confidential", "Page 1 of 2 / Page 2 of 2"),
            VERTRAG_NOTES),
        "jiance-baogao.en.md": reference(
            "jiance-baogao.docx", "Source: `samples/jiance-baogao.docx` (Chinese, Word).", ref_jiance(), JIANCE_NOTES),
        "anquan-peixun.en.md": reference(
            "anquan-peixun.pptx", "Source: `samples/anquan-peixun.pptx` (Chinese, PowerPoint, 4 slides).", ref_anquan(),
            ANQUAN_NOTES),
        "quexian-qingdan.en.md": reference(
            "quexian-qingdan.xlsx", "Source: `samples/quexian-qingdan.xlsx` (Chinese, Excel, 2 sheets).", ref_quexian(),
            QUEXIAN_NOTES),
    }


# =========================================================================== #
# Main
# =========================================================================== #


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=ROOT / "samples", help="output folder (default: the repo's samples/)")
    out: Path = parser.parse_args().out
    for sub in ("", "demo/zh", "reference"):
        (out / sub).mkdir(parents=True, exist_ok=True)

    def write(name: str, data: bytes | str) -> None:
        path = out / name
        path.write_bytes(data.encode() if isinstance(data, str) else data)
        print(f"  {name:42} {path.stat().st_size:>9,} bytes")

    jobs: list[tuple[str, Callable[[], pymupdf.Document]]] = [("fucha-baogao", build_fucha), ("vertragsauszug", build_vertrag)]
    for name, build in jobs:
        doc = build()
        write(f"{name}.pdf", pdf_bytes(doc))
        write(f"{name}-scanned.pdf", scanned(doc))
    jiance = build_jiance()
    write("jiance-baogao.docx", jiance)
    write("anquan-peixun.pptx", build_anquan())
    write("quexian-qingdan.xlsx", build_quexian())
    write("demo/zh/termbase.csv", termbase_csv())
    write("demo/zh/memory.tmx", memory_tmx())
    write("demo/zh/jiance-baogao.vendor.en.xlf", vendor_xliff(jiance))
    for name, text in references().items():
        write(f"reference/{name}", text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
