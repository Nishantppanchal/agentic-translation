"""The Word export, built from segments: the structure (headings, lists, tables, figures, running header and
footer) is kept, the original page layout is not."""

from __future__ import annotations

import io
import re
from collections import Counter

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt

from ..languages import EAST_ASIAN_FONTS, RTL
from ..text import markdown_grid

_md = MarkdownIt("commonmark").enable("table")
_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_DIGITS = re.compile(r"\d+")


def _text(seg: dict, field: str) -> str:
    return (seg.get(field) if field == "source" else seg.get("target")) or seg.get("source") or ""


def _set_language(doc: Document, lang: str) -> None:
    """Document language (spell-check, hyphenation) and an East Asian font for CJK targets."""
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)
    lang_el.set(qn("w:val"), lang)
    if lang in EAST_ASIAN_FONTS:
        lang_el.set(qn("w:eastAsia"), lang)
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), EAST_ASIAN_FONTS[lang])
    if lang.split("-")[0] in RTL:
        bidi = OxmlElement("w:rtl")
        rpr.append(bidi)


def _inline(paragraph, text: str, bold: bool = False) -> None:
    """Markdown inline (bold, italic, code, links) → runs."""
    tokens = _md.parseInline(text)
    children = tokens[0].children if tokens else []
    state = {"strong": bold, "em": False}
    for t in children or []:
        if t.type == "strong_open":
            state["strong"] = True
        elif t.type == "strong_close":
            state["strong"] = bold
        elif t.type == "em_open":
            state["em"] = True
        elif t.type == "em_close":
            state["em"] = False
        elif t.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif t.type in ("text", "code_inline"):
            run = paragraph.add_run(t.content)
            run.bold, run.italic = state["strong"], state["em"]
            if t.type == "code_inline":
                run.font.name = "Consolas"


def _running_parts(segments: list[dict], kind: str, field: str) -> list[str | tuple[str, str]]:
    """The running header/footer as literal text and Word fields: ("PAGE", "1") / ("NUMPAGES", "2").

    Word repeats one header/footer on every page, so a page number copied from page 1 would print
    "Page 1 of 2" on page 2 as well. Numbers that follow the page number become PAGE fields, and a
    constant number next to them that equals the page count becomes NUMPAGES.
    """
    by_page: dict[int, str] = {}
    for s in segments:
        text = _text(s, field).strip() if s.get("kind") == kind else ""
        if text:
            by_page[s["page"]] = f"{by_page[s['page']]} · {text}" if s["page"] in by_page else text
    if not by_page:
        return []
    total = max(s.get("page") or 1 for s in segments)
    shape = lambda t: _DIGITS.sub("#", t)  # noqa: E731
    common = Counter(shape(t) for t in by_page.values()).most_common(1)[0][0]
    pages = sorted(p for p, t in by_page.items() if shape(t) == common)
    first = by_page[pages[0]]
    columns = list(zip(*(map(int, _DIGITS.findall(by_page[p])) for p in pages)))
    # With a single sample, only a short page-number line ("Seite 1 von 2", "- 2 -") counts.
    page_like = len(pages) > 1 or len(re.findall(r"\w+", _DIGITS.sub("", first))) <= 3
    page_col = next((i for i, col in enumerate(columns) if page_like and list(col) == pages), None)  # one per line
    kinds: list[str | None] = [("PAGE" if i == page_col else None) for i in range(len(columns))]
    if page_col is not None:
        kinds = [k or ("NUMPAGES" if set(col) == {total} else None) for k, col in zip(kinds, columns)]
    parts: list[str | tuple[str, str]] = []
    pos = 0
    for m, k in zip(_DIGITS.finditer(first), kinds):
        if k:
            parts += [first[pos:m.start()], (k, m.group())]
            pos = m.end()
    parts.append(first[pos:])
    return [p for p in parts if p]


def _add_field(paragraph, instr: str, shown: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} \\* MERGEFORMAT ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = shown  # what Word shows until it updates the field
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def _table(doc: Document, grid: list[list[str]], style: str = "Table Grid") -> None:
    if not grid:
        return
    cols = max(len(r) for r in grid)
    table = doc.add_table(rows=len(grid), cols=cols)
    table.style = style
    for r, row in enumerate(grid):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = ""
            _inline(cell.paragraphs[0], row[c] if c < len(row) else "", bold=(r == 0))


def to_docx(segments: list[dict], field: str, lang: str, figures: dict[str, bytes]) -> bytes:
    doc = Document()
    _set_language(doc, lang)
    # Running header/footer → the document's real header/footer, with live page-number fields.
    section = doc.sections[0]
    for kind, part in (("header", section.header), ("footer", section.footer)):
        paragraph = part.paragraphs[0]
        for item in _running_parts(segments, kind, field):
            if isinstance(item, tuple):
                _add_field(paragraph, *item)
            else:
                paragraph.add_run(item)
    for seg in segments:
        text = _text(seg, field)
        kind = seg.get("kind")
        if kind in ("header", "footer"):
            continue
        if kind == "figure":
            image = figures.get(seg.get("image") or "")
            if image:
                doc.add_picture(io.BytesIO(image), width=Inches(5.5))
            if text.strip():
                cap = doc.add_paragraph(style="Caption")
                cap.add_run(text.replace("\n", " · "))
            continue
        if kind == "table":
            _table(doc, markdown_grid(text))
            doc.add_paragraph()
            continue
        if kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(text.strip("`\n"))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        for line in text.splitlines() or [""]:
            if not line.strip():
                continue
            if m := _HEADING.match(line):
                doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            elif m := _BULLET.match(line):
                _inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            elif m := _NUMBERED.match(line):
                _inline(doc.add_paragraph(style="List Number"), m.group(1))
            else:
                _inline(doc.add_paragraph(), line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
