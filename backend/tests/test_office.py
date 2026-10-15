"""Word, PowerPoint and Excel files: read from the file (no OCR), translation written back into a copy that keeps
the layout, formatting, objects (footnote marks, pictures, fields) and everything we don't touch."""

from __future__ import annotations

import asyncio
import io
import re
import warnings
import zipfile

import openpyxl
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from translator import office
from translator.office import markup
from translator.office.markup import Piece

from .fixtures_office import (
    docx_full,
    docx_only_pictures,
    docx_pictures,
    labelled_png,
    pptx_deck,
    pptx_pictures,
    xlsx_book,
    xlsx_pictures,
)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FILES = {"docx": docx_full, "pptx": pptx_deck, "xlsx": xlsx_book}


def _pseudo(text: str) -> str:
    """A stand-in translation that keeps the markers: upper case, and a mark at the end."""
    return text.upper() + " ++"


@pytest.mark.parametrize("fmt", list(FILES))
def test_writing_the_source_back_gives_the_same_file_to_read(fmt):
    data = FILES[fmt]()
    segs, _ = office.extract(data, fmt)
    out, report = office.write(data, fmt, [{**s, "target": s["source"]} for s in segs])
    assert not report.mismatch and not report.anchors_moved and not report.formatting_simplified
    again, _ = office.extract(out, fmt)
    key = lambda ss: [(s["id"], s["source"], s["where"], s["kind"]) for s in ss]  # noqa: E731
    assert key(again) == key(segs)


def test_word_pseudo_translation_keeps_formatting_links_objects_and_notes():
    data = docx_full()
    segs, info = office.extract(data, "docx")
    assert info["tracked_changes"] == 2
    out, report = office.write(data, "docx", [{**s, "target": _pseudo(s["source"]) if s["translatable"] else None} for s in segs])
    assert not report.mismatch and not report.anchors_moved
    doc = Document(io.BytesIO(out))
    body = doc.element.body
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    para = next(p for p in doc.paragraphs if "KABELPRITSCHE" in p.text)
    bold = next(r for r in para.runs if r.text == "KABELPRITSCHE")
    assert bold.bold and next(r for r in para.runs if r.text == "NICHT").italic
    assert "\t" in para.text  # the tab survived
    link = body.find(f".//{W}hyperlink")
    assert link.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id") == "rIdLink1"
    assert "".join(t.text for t in link.iter(f"{W}t")) == "WEBSITE DES NETZBETREIBERS"
    note_para = next(p for p in body.iter(f"{W}p") if p.find(f".//{W}footnoteReference") is not None)
    texts = [(e.tag.split("}")[1], e.text) for e in note_para.iter() if e.tag in (f"{W}t", f"{W}footnoteReference")]
    assert texts == [("t", "DIE FRIST"), ("footnoteReference", None), ("t", " IST VERBINDLICH. ++")]
    assert "GEMÄSS VOB/B § 12. ++" in zipfile.ZipFile(io.BytesIO(out)).read("word/footnotes.xml").decode()
    assert "REF _Ref1" in xml and xml.count("HINWEIS IM TEXTFELD ++") == 2  # both copies of the text box
    assert "ERGEBNIS DER PRÜFUNG" in xml  # a field's result is translated, its instruction kept
    assert "<w:ins " not in xml and "<w:del " not in xml and "VOLLSTÄNDIG" in xml
    assert xml.count("<w:drawing>") == docx_xml(data).count("<w:drawing>")
    assert re.search(r'<w:lang w:val="en-GB"', zipfile.ZipFile(io.BytesIO(out)).read("word/styles.xml").decode())
    footer = zipfile.ZipFile(io.BytesIO(out)).read("word/footer1.xml").decode()
    assert "SEITE" in footer and "PAGE" in footer
    header = zipfile.ZipFile(io.BytesIO(out)).read("word/header1.xml").decode()
    assert "VERTRAULICH" in header


def docx_xml(data: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()


def test_powerpoint_pseudo_translation_keeps_breaks_groups_notes_and_shrinks_what_grew():
    data = pptx_deck()
    segs, _ = office.extract(data, "pptx")
    targets = {s["id"]: _pseudo(s["source"]) for s in segs if s["translatable"]}
    targets[next(s["id"] for s in segs if s["source"] == "Kurz")] = "A much longer English text that grows"
    out, report = office.write(data, "pptx", [{**s, "target": targets.get(s["id"])} for s in segs])
    assert not report.mismatch
    prs = Presentation(io.BytesIO(out))
    slide = prs.slides[0]
    assert slide.shapes.title.text == "# ABNAHME UMSPANNWERK NORD ++".removeprefix("# ")
    body = slide.placeholders[1].text_frame
    assert body.paragraphs[0].text.split("\v") == ["ERSTE ZEILE", "NACH DEM UMBRUCH ++"]  # the line break kept
    assert [r.font.bold for r in body.paragraphs[1].runs][:1] == [True]
    texts = [sh.text_frame.text for g in slide.shapes if g.shape_type == 6 for sh in g.shapes]
    assert texts == ["IN DER GRUPPE ++"]
    assert slide.notes_slide.notes_text_frame.text == "SPRECHERNOTIZ ZUR ABNAHME ++"
    assert any(a["where"].endswith("Rectangle 6") and a["scale"] < 1 for a in report.autofit)
    assert 'lang="en-GB"' in zipfile.ZipFile(io.BytesIO(out)).read("ppt/slides/slide1.xml").decode()


def test_excel_keeps_the_file_and_updates_formula_literals():
    data = xlsx_book()
    segs, _ = office.extract(data, "xlsx")
    tr = {"offen": "open", "Beschreibung": "Description", "**Wichtig:** Erdung prüfen": "**Important:** check the earthing"}
    out, report = office.write(data, "xlsx", [{**s, "target": tr.get(s["source"], _pseudo(s["source"]))} for s in segs])
    assert report.formulas_updated == 1
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(out), rich_text=True)
    ws = wb["Mängel"]
    assert (ws["A1"].value, ws["A4"].value) == ("Description", "Description")  # one shared string, both cells
    assert ws["D2"].value == ws["D3"].value == "open" and ws["E1"].value == '=COUNTIF(D:D,"open")'
    assert ws["C2"].value == "=LEN(A2)" and ws["B4"].value == 42
    rich = ws["A3"].value
    assert [(b.font.b if hasattr(b, "font") else None, str(b)) for b in rich][0] == (True, "Important:")
    assert wb.sheetnames == ["Mängel", "Notizen"]
    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(out))
    assert before.namelist() == after.namelist()
    changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
    assert changed <= {"xl/sharedStrings.xml", "xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    assert b"sparkline" in after.read("xl/worksheets/sheet1.xml")


# --------------------------------------------------------------------------- #
# The writer is careful
# --------------------------------------------------------------------------- #


def test_a_source_that_no_longer_matches_is_left_alone():
    data = docx_full()
    segs, _ = office.extract(data, "docx")
    segs = [{**s, "target": "X"} if s["source"] == "# Abnahmeprotokoll" else s for s in segs]
    segs = [{**s, "source": "# Etwas anderes"} if s["target"] == "X" else s for s in segs]
    out, report = office.write(data, "docx", segs)
    assert report.mismatch and "Abnahmeprotokoll" in docx_xml(out)


def test_lost_markers_are_reported_and_the_object_is_kept():
    data = docx_full()
    segs, _ = office.extract(data, "docx")
    targets = {"Die Frist[^1] ist verbindlich.": "The deadline is binding.",
               "Die **Kabelpritsche** im Kellergeschoss ist *nicht* geerdet.\tFrist: 30.04.2026": "The cable tray is not earthed.\tDeadline: 30/04/2026"}
    out, report = office.write(data, "docx", [{**s, "target": targets.get(s["source"])} for s in segs])
    assert len(report.anchors_moved) == 1 and len(report.formatting_simplified) == 0
    assert docx_xml(out).count("footnoteReference") == 1  # the footnote mark is still there


# --------------------------------------------------------------------------- #
# Markup
# --------------------------------------------------------------------------- #


def test_markup_escapes_literal_markdown_and_marks_differences_only():
    base = Piece(text="a * b [x] ", style="")
    bold = Piece(text="fett", style="b", bold=True, payload="B")
    enc = markup.encode([base, bold, Piece(text=" ", style="")])
    assert enc.text == r"a \* b \[x\] **fett** "
    dec = markup.decode(r"a \* b \[x\] **bold** ", enc)
    assert [(o.text, o.like.payload if o.like else None) for o in dec.out] == [("a * b [x] ", None), ("bold", "B"), (" ", None)]


def test_markup_bold_and_italic_together_and_a_split_bold_span():
    enc = markup.encode([Piece(text="Plain ", style=""), Piece(text="both", style="bi", bold=True, italic=True, payload="BI")])
    assert enc.text == "Plain ***both***"
    dec = markup.decode("***beide \\* zusammen***", enc)
    assert [(o.text, o.like.payload) for o in dec.out] == [("beide * zusammen", "BI")]  # one span, one stretch


# --------------------------------------------------------------------------- #
# API and workflow
# --------------------------------------------------------------------------- #


def test_old_and_broken_files_are_refused_with_a_reason():
    from translator.api.routes import documents

    app = FastAPI()
    app.include_router(documents.router)
    client = TestClient(app)
    r = client.post("/api/sessions", files={"file": ("bericht.doc", b"\xd0\xcf\x11\xe0junk", "application/msword")})
    assert r.status_code == 400 and "save it as .docx" in r.json()["detail"]
    r = client.post("/api/sessions", files={"file": ("bericht.docx", b"not a zip", "application/octet-stream")})
    assert r.status_code == 400 and "isn't a readable Word document" in r.json()["detail"]
    with pytest.raises(ValueError, match="password-protected"):
        office.check(b"\xd0\xcf\x11\xe0" + b"\0" * 100, "xlsx")


def test_a_word_document_goes_through_the_pipeline_without_ocr(monkeypatch):
    from google.adk import Runner
    from google.adk.apps import App
    from google.adk.artifacts import InMemoryArtifactService
    from google.adk.sessions import InMemorySessionService
    from google.genai import types

    from translator.agents import pipeline, workers
    from translator.document import state as st
    from translator.exports import available, build

    from .test_flow import _reply, fake_review

    prompts: list[str] = []

    def translate(callback_context, llm_request):
        text = "\n".join(p.text or "" for c in llm_request.contents for p in c.parts or [] if p.text)
        prompts.append(text)
        segs = re.findall(r'<seg id="([^"]+)"[^>]*>\n(.*?)\n</seg>', text, re.S)
        return _reply("\n".join(f'<seg id="{sid}">{_pseudo(src)}</seg>' for sid, src in segs))

    monkeypatch.setattr(workers.translator, "before_model_callback", translate)
    monkeypatch.setattr(workers.reviewer, "before_model_callback", fake_review)

    async def main():
        sessions, artifacts = InMemorySessionService(), InMemoryArtifactService()
        runner = Runner(app=App(name="t", root_agent=pipeline.ingest_workflow), session_service=sessions, artifact_service=artifacts)
        s = await sessions.create_session(app_name="t", user_id="u", state={
            "doc:filename": "bericht.docx", "doc:source_artifact": "source.docx", "doc:source_lang": "German",
            "doc:target_lang": "English", "doc:questions": [], "doc:versions": []})
        await artifacts.save_artifact(app_name="t", user_id="u", session_id=s.id, filename="source.docx",
                                      artifact=types.Part.from_bytes(data=docx_full(), mime_type=office.media_type("x.docx")))
        async for _ in runner.run_async(user_id="u", session_id=s.id, new_message=types.Content(role="user", parts=[types.Part(text="go")])):
            pass
        s = await sessions.get_session(app_name="t", user_id="u", session_id=s.id)
        assert s.state["doc:status"] == "ready" and s.state["doc:page_labels"] == ["Part 1", "Notes"]
        segs = st.get_segments(s.state)
        assert all(x.get("loc") for x in segs) and segs[0]["where"] == "Header"
        assert any('where="Table 1 · R1C2"' in p for p in prompts)  # where each segment is, as context
        assert s.state["doc:versions"][0]["version"] == 1
        assert [f["id"] for f in available(s)][:2] == ["office", "md"]
        content, media, name = await build(artifacts, s, "office", None)
        assert name.endswith(".docx") and media.endswith("wordprocessingml.document")
        assert "KABELPRITSCHE" in docx_xml(content)

    asyncio.run(main())


def test_the_layout_view_gets_the_file_with_bookmarks_and_sheets_as_grids(monkeypatch):
    from google.adk.artifacts import InMemoryArtifactService
    from google.adk.sessions import InMemorySessionService
    from google.genai import types

    from translator.api import services
    from translator.api.routes import office as office_routes
    from translator.document import state as st

    sessions, artifacts = InMemorySessionService(), InMemoryArtifactService()
    monkeypatch.setattr(services, "session_service", sessions)
    monkeypatch.setattr(services, "artifact_service", artifacts)
    monkeypatch.setattr(office_routes, "artifact_service", artifacts)
    office_routes._cache.clear()

    async def make(name: str, data: bytes) -> str:
        segs, info = office.extract(data, office.detect(name))
        segs = [{**x, "target": _pseudo(x["source"]) if x["translatable"] else x["target"]} for x in segs]
        s = await sessions.create_session(app_name=services.APP_NAME, user_id=services.USER, state={
            "doc:filename": name, "doc:source_artifact": "source" + name[-5:], "doc:seg_ids": [x["id"] for x in segs],
            "doc:office": info, **{st.seg_key(x["id"]): x for x in segs}})
        await artifacts.save_artifact(app_name=services.APP_NAME, user_id=services.USER, session_id=s.id, filename="source" + name[-5:],
                                      artifact=types.Part.from_bytes(data=data, mime_type=office.media_type(name)))
        return s.id

    word_id = asyncio.run(make("bericht.docx", docx_full()))
    book_id = asyncio.run(make("maengel.xlsx", xlsx_book()))
    app = FastAPI()
    app.include_router(office_routes.router)
    client = TestClient(app)
    marked = client.get(f"/api/sessions/{word_id}/office/file", params={"side": "target", "marks": True})
    assert marked.status_code == 200 and 'w:name="_seg_p1_2"' in docx_xml(marked.content) and "ABNAHMEPROTOKOLL" in docx_xml(marked.content)
    original = client.get(f"/api/sessions/{word_id}/office/file", params={"side": "source"}).content
    assert original == docx_full() or "Abnahmeprotokoll" in docx_xml(original)
    report = client.get(f"/api/sessions/{word_id}/office/report").json()
    assert report["extract"]["tracked_changes"] == 2 and report["write"]["written"] > 10
    grid = client.get(f"/api/sessions/{book_id}/office/sheets").json()
    first = grid["sheets"][0]
    assert first["name"] == "Mängel" and first["rows"][0][0] == {"v": "Beschreibung", "seg": "p1.1"}
    assert first["rows"][1][1]["v"] == "2026" and first["rows"][3][1]["align"] == "right"
    assert client.get(f"/api/sessions/{word_id}/office/sheets").status_code == 404



# --------------------------------------------------------------------------- #
# Pictures
# --------------------------------------------------------------------------- #

PICTURE_FILES = {"docx": docx_pictures, "pptx": pptx_pictures, "xlsx": xlsx_pictures}
ENGLISH = "Circuit diagram\n**Transformer**"  # a translation with markdown: alt text and captions get plain text


def _translated_figures(segs):
    return [{**s, "source": "Schaltbild\nTrafo", "translatable": True, "target": ENGLISH} if s["kind"] == "figure" else s
            for s in segs]


def test_pictures_become_figure_segments_in_reading_order_and_the_rest_is_counted():
    segs, info = office.extract(docx_pictures(), "docx")
    assert [(s["kind"], s["where"]) for s in segs] == [
        ("figure", "Header · picture 1"), ("paragraph", "Body"), ("alt", "Body · picture 1 · alt text"),
        ("figure", "Body · picture 1 (+1 more)"), ("paragraph", "Body"), ("paragraph", "Body"),
        ("alt", "Body · picture 5 · alt text")]
    alt, figure, emf_alt = segs[2], segs[3], segs[6]
    assert figure["loc"] == "/word/document.xml#1/pic0" and figure["locs"] == [figure["loc"], "/word/document.xml#3/pic0"]
    assert figure["image"] == "figure-p1-2.png" and figure["figure"] == 2 and not figure["translatable"]
    # a picture's own alt text is a segment, even when the picture's text isn't read (an EMF)
    assert (alt["source"], alt["loc"], alt["translatable"]) == ("Übersichtsschaltbild der Anlage", figure["loc"] + "#alt", True)
    assert emf_alt["source"] == "Vektorgrafik des Trafos"
    # the repeated diagram is one segment (read once); icon, decorative and EMF pictures are left out and counted
    assert info["pictures"] == {"found": 6, "segments": 2, "alt_texts": 2,
                                "skipped": {"small": 1, "decorative": 1, "vector": 1}}
    assert info["not_translated"] == {"pictures in EMF/WMF format (not read)": 1}
    assert set(info["picture_files"]) == {s["sha1"] for s in segs if s["kind"] == "figure"}
    segs, _ = office.extract(pptx_pictures(), "pptx")  # python-pptx's alt text is the file name: not a segment
    assert [(s["kind"], s["where"]) for s in segs] == [
        ("heading", "Slide 1 · Title"), ("figure", "Slide 1 · picture 1 (+1 more)"), ("note", "Slide 1 · Notes"),
        ("heading", "Slide 2 · Title")]
    segs, info = office.extract(xlsx_pictures(), "xlsx")
    assert [(s["kind"], s["source"], s["where"]) for s in segs if s["kind"] in ("alt", "figure")] == [
        ("alt", "Schaltbild", "Mängel · picture at D2 · alt text"), ("figure", "", "Mängel · picture at D2")]
    assert info["pictures"]["skipped"] == {"decorative": 1}
    segs, info = office.extract(docx_pictures(), "docx", read_pictures=False)
    assert [s["kind"] for s in segs if s["kind"] in ("alt", "figure")] == ["alt", "alt"]
    assert info["pictures"]["skipped"] == {"decorative": 1, "not_read": 5} and not info["picture_files"]


def test_a_file_with_only_pictures_can_be_uploaded():
    office.check(docx_only_pictures(), "docx")
    segs, _ = office.extract(docx_only_pictures(), "docx")
    assert [s["kind"] for s in segs] == ["figure"]


def test_word_pictures_get_alt_text_and_a_caption_in_the_body_only():
    data = docx_pictures()
    segs, _ = office.extract(data, "docx")
    alts = {"Übersichtsschaltbild der Anlage": "Overview diagram of the installation"}  # the EMF's alt text: not translated
    out, report = office.write(data, "docx", [{**s, "target": alts.get(s["source"])} if s["kind"] == "alt" else s
                                              for s in _translated_figures(segs)])
    assert not report.mismatch
    assert report.pictures == {"alt_text": 3, "alt_translated": 1, "captions": 2, "notes": 0}
    xml, header = docx_xml(out), zipfile.ZipFile(io.BytesIO(out)).read("word/header1.xml").decode()
    picture_text = "Text in the picture: Circuit diagram; Transformer"
    # with its own alt text translated, then the picture's text; the repeat (no alt text of its own): the text only
    assert xml.count(f'descr="Overview diagram of the installation — {picture_text}"') == 2  # wp:docPr and pic:cNvPr
    assert xml.count(f'descr="{picture_text}"') == 2
    assert 'descr="Vektorgrafik des Trafos"' in xml  # untranslated and nothing read in it: left as it was
    assert f'descr="{picture_text}"' in header and "Caption" not in header
    assert xml.count("<w:drawing>") == docx_xml(data).count("<w:drawing>")
    doc = Document(io.BytesIO(out))
    texts = [(p.style.name, p.text) for p in doc.paragraphs]
    assert texts[1:3] == [("Normal", ""), ("Caption", "Circuit diagram\nTransformer")]  # right under the picture
    assert sum(t == "Circuit diagram\nTransformer" for _, t in texts) == 2
    caption = next(p for p in doc.paragraphs if p.style.name == "Caption")
    assert caption.runs[0]._r.find(f"{W}rPr/{W}lang").get(f"{W}val") == "en-GB"


def test_powerpoint_pictures_get_alt_text_and_a_line_in_the_slide_notes():
    data = pptx_pictures()
    segs, _ = office.extract(data, "pptx")
    out, report = office.write(data, "pptx", _translated_figures(segs))
    assert not report.mismatch and report.pictures == {"alt_text": 2, "alt_translated": 0, "captions": 0, "notes": 2}
    prs = Presentation(io.BytesIO(out))
    first, second = prs.slides
    assert first.notes_slide.notes_text_frame.text == (
        "Sprechernotiz\n\nText in pictures (English):\nPicture 1: Circuit diagram\vTransformer")
    assert second.has_notes_slide and second.notes_slide.notes_text_frame.text.startswith("Text in pictures (English):")
    pics = [sh for sh in first.shapes if sh.shape_type == 13]
    assert pics[0]._element.nvPicPr.cNvPr.get("descr") == "Text in the picture: Circuit diagram; Transformer"
    assert pics[1]._element.nvPicPr.cNvPr.get("descr") == "image.png"  # the icon was left alone


def test_excel_pictures_get_alt_text_and_nothing_else_changes():
    data = xlsx_pictures()
    segs, _ = office.extract(data, "xlsx")
    out, report = office.write(data, "xlsx", _translated_figures(segs))  # its alt text "Schaltbild" not translated
    assert not report.mismatch and report.pictures["alt_text"] == 1 and report.pictures["alt_translated"] == 0
    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(out))
    assert before.namelist() == after.namelist()
    changed = {n for n in before.namelist() if before.read(n) != after.read(n)}
    assert changed <= {"xl/drawings/drawing1.xml", "xl/sharedStrings.xml"} and "xl/drawings/drawing1.xml" in changed
    drawing = after.read("xl/drawings/drawing1.xml").decode()
    assert 'descr="Schaltbild — Text in the picture: Circuit diagram; Transformer"' in drawing  # kept, then the text
    translated = [{**s, "target": "Circuit diagram"} if s["kind"] == "alt" else s for s in _translated_figures(segs)]
    drawing = zipfile.ZipFile(io.BytesIO(office.write(data, "xlsx", translated)[0])).read("xl/drawings/drawing1.xml").decode()
    assert 'descr="Circuit diagram — Text in the picture: Circuit diagram; Transformer"' in drawing


def test_an_alt_text_alone_is_translated_when_nothing_was_read_in_the_picture():
    data = docx_pictures()
    segs, _ = office.extract(data, "docx", read_pictures=False)
    segs = [{**s, "target": "Vector drawing of the transformer"} if s["source"] == "Vektorgrafik des Trafos" else s for s in segs]
    out, report = office.write(data, "docx", segs)
    assert 'descr="Vector drawing of the transformer"' in docx_xml(out) and report.pictures["captions"] == 0
    assert report.pictures["alt_translated"] == 1 and not report.mismatch


def test_a_picture_that_changed_in_the_file_is_left_alone():
    data = pptx_pictures()
    segs, _ = office.extract(data, "pptx")
    segs = [{**s, "sha1": "0" * 40} if s["kind"] == "figure" else s for s in _translated_figures(segs)]
    out, report = office.write(data, "pptx", segs)
    assert report.mismatch == ["p1.2"] and report.pictures["notes"] == 0


def _picture_reader(monkeypatch, texts: dict[bytes, str]) -> list[bytes]:
    """The figure reader as a fake vision model: the text "in" each picture, by the PNG it is shown."""
    from translator.agents import workers

    from .test_flow import _reply

    seen: list[bytes] = []

    def look(callback_context, llm_request):
        png = next(p.inline_data.data for c in llm_request.contents for p in c.parts or [] if p.inline_data)
        seen.append(png)
        return _reply(f"<kind>picture</kind>\n<text>\n{texts.get(png, '')}\n</text>")

    monkeypatch.setattr(workers.figure_reader, "before_model_callback", look)
    return seen


def test_word_pictures_go_through_the_pipeline_read_once_and_cached(monkeypatch):
    from google.adk import Runner
    from google.adk.apps import App
    from google.adk.artifacts import InMemoryArtifactService
    from google.adk.sessions import InMemorySessionService
    from google.genai import types

    from translator.agents import pipeline, workers
    from translator.document import state as st
    from translator.exports import build
    from translator.office.pictures import to_png

    from .test_flow import _reply, fake_review

    texts = {to_png(labelled_png()): "Schaltbild\nTrafo", to_png(labelled_png("Firma", color=(1, 0.9, 0.8))): "Firma GmbH"}
    seen = _picture_reader(monkeypatch, texts)

    def translate(callback_context, llm_request):
        text = "\n".join(p.text or "" for c in llm_request.contents for p in c.parts or [] if p.text)
        segs = re.findall(r'<seg id="([^"]+)"[^>]*>\n(.*?)\n</seg>', text, re.S)
        return _reply("\n".join(f'<seg id="{sid}">{_pseudo(src)}</seg>' for sid, src in segs))

    monkeypatch.setattr(workers.translator, "before_model_callback", translate)
    monkeypatch.setattr(workers.reviewer, "before_model_callback", fake_review)

    async def main():
        sessions, artifacts = InMemorySessionService(), InMemoryArtifactService()
        runner = Runner(app=App(name="t", root_agent=pipeline.ingest_workflow), session_service=sessions, artifact_service=artifacts)
        s = await sessions.create_session(app_name="t", user_id="u", state={
            "doc:filename": "bilder.docx", "doc:source_artifact": "source.docx", "doc:source_lang": "German",
            "doc:target_lang": "English", "doc:questions": [], "doc:versions": []})
        await artifacts.save_artifact(app_name="t", user_id="u", session_id=s.id, filename="source.docx",
                                      artifact=types.Part.from_bytes(data=docx_pictures(), mime_type=office.media_type("x.docx")))

        async def run():
            async for _ in runner.run_async(user_id="u", session_id=s.id,
                                            new_message=types.Content(role="user", parts=[types.Part(text="go")])):
                pass
            return await sessions.get_session(app_name="t", user_id="u", session_id=s.id)

        done = await run()
        assert done.state["doc:status"] == "ready" and len(seen) == 2  # two different pictures, each read once
        figures = [x for x in st.get_segments(done.state) if x["kind"] == "figure"]
        assert [(f["source"], f["target"]) for f in figures] == [
            ("Firma GmbH", "FIRMA GMBH ++"), ("Schaltbild\nTrafo", "SCHALTBILD\nTRAFO ++")]
        image = await artifacts.load_artifact(app_name="t", user_id="u", session_id=s.id, filename=figures[1]["image"])
        assert image.inline_data.data == to_png(labelled_png())
        assert "picture_files" not in done.state["doc:office"] and done.state["doc:office"]["pictures"]["segments"] == 2
        content, _, _ = await build(artifacts, done, "office", None)
        xml = docx_xml(content)
        assert 'descr="ÜBERSICHTSSCHALTBILD DER ANLAGE ++ — Text in the picture: SCHALTBILD; TRAFO ++"' in xml
        assert 'descr="Text in the picture: SCHALTBILD; TRAFO ++"' in xml  # the repeat: no alt text of its own
        assert 'descr="VEKTORGRAFIK DES TRAFOS ++"' in xml  # the EMF's alt text, translated
        assert [x["kind"] for x in st.get_segments(done.state)].count("alt") == 2
        assert any("Read the text in 2 pictures" in (e.content.parts[0].text or "") for e in done.events
                   if e.content and e.content.parts)
        again = await run()  # a retry reads the file again, not the pictures
        assert again.state["doc:status"] == "ready" and len(seen) == 2

    asyncio.run(main())


def test_the_chat_says_why_pictures_were_not_read():
    from translator.agents.steps import extract

    assert extract._pictures_line(0, 0, {"pictures": {"found": 3, "alt_texts": 1}}, "the reading model doesn't read images") == (
        "The text in the file's 3 pictures wasn't read: the reading model doesn't read images. "
        "The alt text of 1 picture (what its author wrote for it) is translated too.")
